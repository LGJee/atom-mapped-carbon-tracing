"""Build a small synthetic input to exercise transport and enzymatic mapping."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment


def main():
    root = Path(__file__).resolve().parent
    workbook = Workbook()
    reactions = workbook.active
    reactions.title = "Nonzero"
    reactions.append(["ReactionID", "ReactionName", "Flux", "Equation"])
    reactions.append(
        [
            "example_transport",
            "Acetaldehyde transport",
            1,
            "acetaldehyde[e] => acetaldehyde[c]",
        ]
    )
    reactions.append(
        [
            "example_reduction",
            "Acetaldehyde reduction",
            1,
            "acetaldehyde[c] + NADH[c] + H+[c] => ethanol[c] + NAD(+)[c]",
        ]
    )
    metabolites = workbook.create_sheet("Metabolites")
    metabolites.append(
        ["MetaboliteID", "MetaboliteName", "SMILES_Yeast9", "SMILES_Manual"]
    )
    nadh = "NC(=O)C1=CN([C@@H]2O[C@H](COP(=O)(O)OP(=O)(O)OC[C@H]3O[C@@H](n4cnc5c(N)ncnc54)[C@H](O)[C@@H]3O)[C@@H](O)[C@H]2O)C=CC1"
    nad = "NC(=O)c1ccc[n+]([C@@H]2O[C@H](COP(=O)(O)OP(=O)(O)OC[C@H]3O[C@@H](n4cnc5c(N)ncnc54)[C@H](O)[C@@H]3O)[C@@H](O)[C@H]2O)c1"
    for number, (name, smiles) in enumerate(
        (
            ("acetaldehyde[e]", "CC=O"),
            ("acetaldehyde[c]", "CC=O"),
            ("NADH[c]", nadh),
            ("H+[c]", "[H+]"),
            ("ethanol[c]", "CCO"),
            ("NAD(+)[c]", nad),
        ),
        1,
    ):
        metabolites.append(["example_" + str(number), name, smiles, ""])
    for sheet, widths in (
        (reactions, (25, 42, 14, 50)),
        (metabolites, (18, 38, 65, 65)),
    ):
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[sheet.cell(1, column).column_letter].width = width
    for row in reactions.iter_rows(min_col=4, max_col=4):
        row[0].alignment = Alignment(wrap_text=True)
    workbook.save(root / "Example_1_input.xlsx")
    workbook.close()
    print(
        "Created examples/Example_1_input.xlsx (synthetic fluxes; not a Yeast9 pFBA solution)."
    )


if __name__ == "__main__":
    main()
