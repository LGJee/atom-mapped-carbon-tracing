"""Check script settings, CLI precedence and the independent propagation entry."""

import sys
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from unittest.mock import patch

from carbon_flow_tools.cli import discover_inputs, main
from carbon_flow_tools.excel_io import input_sheets
from scripts.test_parallel_carbon import make_input
import step2_atom_mapping as mapping
import step3_carbon_flow as tracing


def run():
    with patch.object(mapping, "BATCH_SIZE", 7), patch.object(
        mapping, "MAX_REACTION_ATOMS", 500
    ):
        settings = mapping.settings()
    settings.pop("analysis_plan")
    settings["python_overrides"] = {"LocalMapper": sys.executable}
    assert "rounds" not in settings
    assert "models" not in settings
    with tempfile.TemporaryDirectory() as temporary:
        folder = Path(temporary)
        settings["input_dir"] = str(folder)
        try:
            discover_inputs(folder)
        except FileNotFoundError as error:
            assert "Run step1" in str(error)
        else:
            raise AssertionError("Empty input folder accepted")
        make_input(folder / "non_IUPD_1_input.xlsx", "B[c]")
        custom = folder / "Mixed_Culture_1_custom.xlsx"
        make_input(custom, "B[c]")
        workbook = load_workbook(custom)
        workbook["Nonzero"].title = "My reactions"
        workbook["Metabolites"].title = "My structures"
        workbook["My structures"].delete_cols(3, 2)
        workbook.save(custom)
        workbook.close()
        for name in (
            "non_IUPD_2_atom_mapping.xlsx",
            "non_IUPD_3_carbon_flow.xlsx",
            "~$non_IUPD_1_input.xlsx",
            "non_IUPD_1_input.tmp.xlsx",
        ):
            (folder / name).touch()
        assert set(discover_inputs(folder)) == {"non_IUPD", "Mixed_Culture"}
        with patch("carbon_flow_tools.pipeline.run_analyses") as predict:
            assert main("map", settings, ["non_IUPD", "--mapper", "RXNMapper"]) == 0
            assert predict.call_args[0][2:4] == (7, 500)
            assert predict.call_args[0][5] == ["RXNMapper"]
            assert list(predict.call_args[0][0]) == ["non_IUPD"]
            assert main("map", settings, []) == 0
            assert (
                predict.call_args[0][0]["Mixed_Culture"]["reaction_sheet"]
                == "My reactions"
            )
            assert (
                predict.call_args[0][0]["Mixed_Culture"]["metabolite_sheet"]
                == "My structures"
            )
        duplicate = folder / "non_IUPD_1_duplicate.xlsx"
        duplicate.touch()
        try:
            discover_inputs(folder)
        except ValueError as error:
            assert "Multiple step 1 inputs for model non_IUPD" in str(error)
        else:
            raise AssertionError("Duplicate model accepted")
        duplicate.unlink()
        workbook = load_workbook(custom)
        workbook.copy_worksheet(workbook["My reactions"])
        workbook.save(custom)
        workbook.close()
        with pd.ExcelFile(custom, engine="openpyxl") as excel:
            try:
                input_sheets(excel)
            except ValueError as error:
                assert "exactly one reaction table" in str(error)
            else:
                raise AssertionError("Ambiguous reaction sheets accepted")
        settings["input_dir"] = str(folder / "missing")
        with patch("carbon_flow_tools.pipeline.run_analyses") as predict:
            assert main("map", settings, ["--example"]) == 0
            assert list(predict.call_args[0][0]) == ["Example"]
    assert tracing.INPUT_DIR == mapping.OUTPUT_DIR
    assert tracing.TRACING_PYTHON == ".envs/localmapper/bin/python"
    assert tracing.settings()["allow_inorganic_carbon_recycling"] is False
    assert tracing.settings()["export_exclusive_source_metabolites"] is False
    with patch.object(tracing, "EXPORT_EXCLUSIVE_SOURCE_METABOLITES", True):
        assert tracing.settings()["export_exclusive_source_metabolites"] is True

    with patch.object(tracing, "ALLOW_INORGANIC_CARBON_RECYCLING", True):
        assert tracing.settings()["allow_inorganic_carbon_recycling"] is True
    with patch.object(tracing, "ROUNDS", 6), patch.object(
        tracing, "MAX_CARBON_PROCESSES", 2
    ):
        settings = tracing.settings()
    settings.pop("analysis_plan")
    settings["tracing_python"] = sys.executable
    settings["carbon_sources"] = {
        "non_IUPD": ["D-glucose[e]"],
        "IUPD": ["D-glucose[e]", "isoprenol[e]"],
    }
    assert "mappers" not in settings and "batch_size" not in settings
    assert "analysis_methods" not in settings
    assert "models" not in settings
    with tempfile.TemporaryDirectory() as temporary:
        folder = Path(temporary)
        settings["input_dir"] = str(folder)
        try:
            discover_inputs(folder, "trace")
        except FileNotFoundError as error:
            assert "Run step2_atom_mapping.py" in str(error)
        else:
            raise AssertionError("Empty mapping folder accepted")
        for name in (
            "non_IUPD_2_atom_mapping.xlsx",
            "non_IUPD_2_atom_mapping_AIGenerated.xlsx",
            "Ignore_2_atom_mapping_backup.xlsx",
            "IUPD_2_atom_mapping.xlsx",
            "~$non_IUPD_2_atom_mapping.xlsx",
            "non_IUPD_2_atom_mapping.tmp.xlsx",
            "non_IUPD_1_input.xlsx",
            "non_IUPD_3_carbon_flow.xlsx",
        ):
            (folder / name).touch()
        assert set(discover_inputs(folder, "trace")) == {"non_IUPD", "IUPD"}
        with patch(
            "carbon_flow_tools.cli.build_mapper_configs",
            side_effect=AssertionError("Trace requested mapper settings"),
        ), patch("carbon_flow_tools.pipeline.run_carbon_flow") as propagate:
            assert main("trace", settings, ["non_IUPD", "--rounds", "20"]) == 0
            assert propagate.call_args[0][1] == 20
            job = propagate.call_args[0][0]["non_IUPD"]
            assert job["mapping_file"] == folder / "non_IUPD_2_atom_mapping.xlsx"
            assert job["mapping_files"] == [
                folder / "non_IUPD_2_atom_mapping.xlsx",
                folder / "non_IUPD_2_atom_mapping_AIGenerated.xlsx",
            ]
            assert job["carbon_sources"] == ["D-glucose[e]"]
            assert job["allow_inorganic_carbon_recycling"] is False
            assert job["export_exclusive_source_metabolites"] is False
            settings["export_exclusive_source_metabolites"] = True
            assert "input_file" not in job
            assert main("trace", settings, []) == 0
            assert propagate.call_args[0][1] == 6
            assert propagate.call_args[0][0]["IUPD"]["export_exclusive_source_metabolites"] is True
            assert propagate.call_args[1]["max_processes"] == 2
            assert propagate.call_args[0][0]["IUPD"]["carbon_sources"] == [
                "D-glucose[e]",
                "isoprenol[e]",
            ]
            custom = folder / "Mixed_Culture_2_atom_mapping.xlsx"
            custom.touch()
            # Selecting non-IUPD should not require sources for an unrelated new model.
            assert main("trace", settings, ["non_IUPD"]) == 0
            before = propagate.call_count
            try:
                main("trace", settings, ["Mixed_Culture"])
            except SystemExit as error:
                assert error.code == 2
            else:
                raise AssertionError("Unconfigured carbon sources accepted")
            assert propagate.call_count == before
            settings["carbon_sources"] = {"Mixed_Culture": ["A[e]", "B[e]"]}
            assert main("trace", settings, ["Mixed_Culture"]) == 0
            job = propagate.call_args[0][0]["Mixed_Culture"]
            assert job["mapping_file"] == custom
            assert job["carbon_sources"] == ["A[e]", "B[e]"]
            settings["allow_inorganic_carbon_recycling"] = True
            assert main("trace", settings, ["Mixed_Culture"]) == 0
            assert (
                propagate.call_args[0][0]["Mixed_Culture"][
                    "allow_inorganic_carbon_recycling"
                ]
                is True
            )
        settings["input_dir"] = str(folder / "missing")
        with patch("carbon_flow_tools.pipeline.run_carbon_flow") as propagate:
            assert main("trace", settings, ["--example"]) == 0
            assert propagate.call_args[0][1] == 1
            assert list(propagate.call_args[0][0]) == ["Example"]
    print(
        "PASS: automatic stage inputs, custom sources/order, missing-source protection and script overrides"
    )


if __name__ == "__main__":
    run()
