"""Focused COBRApy preparation checks; no atom-mapping models are loaded."""

import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from cobra import Model, Metabolite, Reaction
from cobra.exceptions import OptimizationError
from cobra.flux_analysis import pfba
import cobra.io
from openpyxl import load_workbook

from carbon_flow_tools.model_preparation import (
    check_solution,
    close_other_carbon_uptake,
    export_model,
    glucose_bound,
    read_model,
    save_plan,
)
from carbon_flow_tools.cli import load_analysis_plan
import step1_prepare_carbon_tracing_inputs as step1


def fixture():
    model = Model("preparation_fixture")
    model.solver = "glpk"
    a = Metabolite("a_e", name="A", formula="C2H4O", charge=0, compartment="e")
    b = Metabolite("b_c", name="B", formula="C2H4O", charge=0, compartment="c")
    exchange = Reaction("EX_a", lower_bound=-10, upper_bound=1000)
    exchange.add_metabolites({a: -1})
    transport = Reaction("transport", lower_bound=0, upper_bound=1000)
    transport.add_metabolites({a: -1, b: 1})
    transport.gene_reaction_rule = "g1"
    growth = Reaction("growth", lower_bound=0, upper_bound=1000)
    growth.add_metabolites({b: -1})
    model.add_reactions([exchange, transport, growth])
    model.objective = growth
    return model


class PreparationTests(unittest.TestCase):
    def test_pfba_reports_biological_flux_separately_from_flux_sum(self):
        model = fixture()
        model.objective = "growth"
        self.assertEqual(glucose_bound(model, "EX_a"), -10)
        solution = pfba(model)
        self.assertAlmostEqual(check_solution(model, solution), 10)
        self.assertAlmostEqual(solution.objective_value, 30)
        with self.assertRaises(KeyError):
            model.reactions.get_by_id("missing_growth_reaction")
        model.reactions.get_by_id("EX_a").lower_bound = 0
        with self.assertRaises(ValueError):
            glucose_bound(model, "EX_a")

    def test_medium_filter_preserves_oxygen_and_both_exchange_signs(self):
        model = fixture()
        oxygen = Metabolite("oxygen_e", formula="O2", compartment="e")
        carbon = Metabolite("carbon_e", formula="C2H4", compartment="e")
        chloride = Metabolite("chloride_e", formula="Cl", compartment="e")
        for key, met, coefficient in (
            ("oxygen", oxygen, -1),
            ("other", carbon, -1),
            ("positive", carbon, 1),
            ("chloride", chloride, -1),
        ):
            reaction = Reaction(key, lower_bound=-3, upper_bound=5)
            reaction.add_metabolites({met: coefficient})
            model.add_reactions([reaction])
        close_other_carbon_uptake(model, "EX_a")
        self.assertEqual(model.reactions.get_by_id("EX_a").bounds, (-10, 1000))
        self.assertEqual(model.reactions.get_by_id("oxygen").bounds, (-3, 5))
        self.assertEqual(model.reactions.get_by_id("chloride").bounds, (-3, 5))
        self.assertEqual(model.reactions.get_by_id("other").bounds, (0, 5))
        self.assertEqual(model.reactions.get_by_id("positive").bounds, (-3, 0))
        model.reactions.get_by_id("other").lower_bound = -3
        close_other_carbon_uptake(model, ["EX_a", "other"])
        self.assertEqual(model.reactions.get_by_id("other").lower_bound, -3)
        with self.assertRaises(KeyError):
            close_other_carbon_uptake(model, ["missing_exchange"])

    def test_gene_knockout_and_infeasible_output_are_checked(self):
        model = fixture()
        with self.assertRaises(KeyError):
            model.genes.get_by_id("missing_gene")
        model.genes.get_by_id("g1").knock_out()
        self.assertFalse(model.genes.get_by_id("g1").functional)
        self.assertEqual(model.reactions.get_by_id("transport").bounds, (0, 0))
        model.reactions.get_by_id("growth").lower_bound = 1
        with self.assertRaises((ValueError, OptimizationError)):
            pfba(model)

    def test_export_rejects_invalid_solutions_before_creating_output(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            model = fixture()
            solution = pfba(model)
            solution.status = "infeasible"
            with self.assertRaisesRegex(ValueError, "finite optimal solution"):
                export_model(model, solution, "Example", ["A[e]"], root)
            self.assertFalse((root / "results").exists())
            solution.status = "optimal"
            solution.fluxes.loc["transport"] = float("nan")
            with self.assertRaisesRegex(ValueError, "finite optimal solution"):
                export_model(model, solution, "Example", ["A[e]"], root)
            self.assertFalse((root / "results").exists())

    def test_native_data_and_portable_json_roundtrip(self):
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            model = fixture()
            for name, writer in (
                ("model.mat", cobra.io.save_matlab_model),
                ("model.json", cobra.io.save_json_model),
                ("model.xml", cobra.io.write_sbml_model),
                ("model.yml", cobra.io.save_yaml_model),
            ):
                path = folder / name
                writer(model, str(path))
                loaded = read_model(path)
                self.assertEqual(
                    [r.id for r in loaded.reactions], [r.id for r in model.reactions]
                )
                self.assertEqual(loaded.reactions.get_by_id("EX_a").bounds, (-10, 1000))
                self.assertEqual(loaded._tracing_source["path"], str(path.resolve()))
                self.assertEqual(
                    loaded._tracing_source["sha256"],
                    hashlib.sha256(path.read_bytes()).hexdigest(),
                )
                self.assertAlmostEqual(check_solution(loaded, pfba(loaded)), 10)
            with self.assertRaises(FileNotFoundError):
                read_model(folder / "absent.mat")

    def test_export_schema_identifiers_sources_font_and_plan(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            model = fixture()
            model.metabolites.get_by_id("a_e").annotation["inchi_key"] = "fixture-key"
            solution = pfba(model)
            entry = export_model(model, solution, "MyCulture", ["A[e]"], root)
            path = root / entry["input_file"]
            self.assertEqual(
                entry["input_sha256"], hashlib.sha256(path.read_bytes()).hexdigest()
            )
            book = load_workbook(path)
            try:
                self.assertEqual(book.sheetnames, ["Nonzero", "Metabolites"])
                self.assertEqual(
                    [c.value for c in book["Nonzero"][1]],
                    ["ReactionID", "ReactionName", "Flux", "Equation"],
                )
                self.assertEqual(book["Nonzero"]["D2"].value, "A[e] <=> ")
                self.assertTrue(
                    all(
                        c.font.name == "Calibri" and c.font.scheme is None
                        for sheet in book
                        for row in sheet
                        for c in row
                        if c.value is not None
                    )
                )
                self.assertIn("InChIKey", [c.value for c in book["Metabolites"][1]])
            finally:
                book.close()
            save_plan(root, {"MyCulture": entry}, "results")
            with patch("carbon_flow_tools.cli.ROOT", root):
                _, jobs = load_analysis_plan(
                    root / "data/analysis_plan.json", "map", {}
                )
            self.assertEqual(list(jobs), ["MyCulture"])
            self.assertEqual(
                json.loads((root / "data/analysis_plan.json").read_text())["models"][
                    "MyCulture"
                ]["carbon_sources"],
                ["A[e]"],
            )
            saved = path.read_bytes()
            for changes in (
                {"sources": []},
                {"sources": ["A[e]", "A[e]"]},
                {"sources": ["absent[c]"]},
                {"tolerance": 0},
                {"output_dir": "../outside"},
            ):
                arguments = dict(name="MyCulture", sources=["A[e]"], project=root)
                arguments.update(changes)
                with self.assertRaises(ValueError):
                    export_model(model, solution, **arguments)
            self.assertEqual(path.read_bytes(), saved)

    def test_failed_preparation_invalidates_previous_plan(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            (root / "data").mkdir()
            plan = root / "data/analysis_plan.json"
            plan.write_text(json.dumps(dict(version=3)), encoding="utf-8")
            with patch.object(step1, "ROOT", root), patch.object(
                step1,
                "prepare_models",
                side_effect=ValueError("Missing configured model"),
            ):
                with self.assertRaisesRegex(ValueError, "Missing configured model"):
                    step1.main()
            self.assertFalse(plan.exists())

    def test_each_model_exports_its_own_ordered_sources(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            first = fixture()
            second = fixture()
            prepared = {
                "FirstCulture": (first, pfba(first), ["A[e]"]),
                "MixedCulture": (second, pfba(second), ["B[c]", "A[e]"]),
            }
            with patch.object(step1, "ROOT", root), patch.object(
                step1, "prepare_models", return_value=prepared
            ):
                step1.main()
            models = json.loads((root / "data/analysis_plan.json").read_text())[
                "models"
            ]
            self.assertEqual(models["FirstCulture"]["carbon_sources"], ["A[e]"])
            self.assertEqual(models["MixedCulture"]["carbon_sources"], ["B[c]", "A[e]"])

    def test_export_records_model_source_without_extra_user_settings(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            folder = root / "data" / "my_species"
            folder.mkdir(parents=True)
            source = folder / "model.json"
            cobra.io.save_json_model(fixture(), str(source))
            loaded = read_model(source)
            entry = export_model(loaded, pfba(loaded), "Example", ["A[e]"], root)
            self.assertEqual(
                entry["source_model"]["path"], "data/my_species/model.json"
            )
            self.assertEqual(
                entry["source_model"]["sha256"],
                hashlib.sha256(source.read_bytes()).hexdigest(),
            )
            save_plan(root, {"Example": entry}, "results")
            with patch("carbon_flow_tools.cli.ROOT", root):
                _, jobs = load_analysis_plan(
                    root / "data/analysis_plan.json", "map", {}
                )
            self.assertEqual(
                jobs["Example"]["source_model"]["path"], str(source.resolve())
            )

    def test_configured_environment_is_preferred_and_missing_dependency_explained(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            python = root / step1.PREPARATION_PYTHON
            python.parent.mkdir(parents=True)
            python.touch()
            with patch.object(step1, "ROOT", root), patch.object(
                step1.os, "execv", side_effect=RuntimeError("interpreter selected")
            ) as launch:
                with self.assertRaisesRegex(RuntimeError, "interpreter selected"):
                    step1.main()
                self.assertEqual(launch.call_args[0][0], str(python))
            python.unlink()
            with patch.object(step1, "ROOT", root), patch.object(
                step1.importlib.util, "find_spec", return_value=None
            ):
                with self.assertRaisesRegex(SystemExit, "setup_environments.sh"):
                    step1.main()

    def test_unused_supplied_carbon_source_is_exported_without_forcing_uptake(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            model = fixture()
            unused = Metabolite("x_e", name="X", formula="C5H10O5", compartment="e")
            exchange = Reaction("EX_x", lower_bound=-5, upper_bound=1000)
            exchange.add_metabolites({unused: -1})
            model.add_reactions([exchange])
            solution = pfba(model)
            self.assertAlmostEqual(solution.fluxes["EX_x"], 0)
            entry = export_model(model, solution, "Mixed", ["A[e]", "X[e]"], root)
            self.assertEqual(entry["carbon_sources"], ["A[e]", "X[e]"])
            record = entry["preparation"]
            self.assertEqual(record["source_uptake"]["X[e]"]["net_import"], 0)
            self.assertEqual(record["source_uptake"]["A[e]"]["net_import"], 10)
            self.assertFalse(record["source_uptake"]["X[e]"]["present_in_nonzero_network"])
            self.assertEqual(record["objective_reactions"], {"growth": 1})
            self.assertEqual(record["objective_value"], 10)
            self.assertEqual(record["total_absolute_flux"], 30)
            book = load_workbook(root / entry["input_file"], read_only=True)
            try:
                self.assertIn("x_e", [r[0] for r in book["Metabolites"].iter_rows(min_row=2, values_only=True)])
                self.assertNotIn("EX_x", [r[0] for r in book["Nonzero"].iter_rows(min_row=2, values_only=True)])
            finally:
                book.close()

    def test_case_plan_and_script_provenance_preserve_default_plan(self):
        from carbon_flow_tools.model_preparation import invalidate_plan
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            default = root / "data/analysis_plan.json"
            default.parent.mkdir()
            default.write_text('{"version":3,"default":true}')
            previous = default.read_bytes()
            script = root / "results/results_by_AI/mixed/step1_prepare_carbon_tracing_inputs.py"
            script.parent.mkdir(parents=True)
            script.write_text("# Case-specific model recipe\n")
            model = fixture()
            entry = export_model(model, pfba(model), "Mixed", ["A[e]"], root,
                                 "results/results_by_AI/mixed")
            path = save_plan(root, {"Mixed": entry}, "results/results_by_AI/mixed",
                             "results/results_by_AI/mixed/analysis_plan.json", script_file=script)
            plan = json.loads(path.read_text())
            self.assertEqual(plan["models"]["Mixed"]["preparation_script"]["sha256"],
                             hashlib.sha256(script.read_bytes()).hexdigest())
            with patch("carbon_flow_tools.cli.ROOT", root):
                settings, jobs = load_analysis_plan(path, "trace", {"rounds": 2})
            self.assertEqual(settings["output_dir"], "results/results_by_AI/mixed")
            self.assertEqual(jobs["Mixed"]["mapping_file"].parent, script.parent)
            invalidate_plan(root, path)
            self.assertFalse(path.exists())
            self.assertEqual(default.read_bytes(), previous)
            unrelated = root / "data/model.json"
            unrelated.write_text('{"metabolites":[]}')
            with self.assertRaises(ValueError):
                invalidate_plan(root, unrelated)
            self.assertTrue(unrelated.exists())

    def test_step1_plan_override_routes_outputs_to_its_case(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            default = root / "data/analysis_plan.json"
            default.parent.mkdir()
            default.write_text('{"version":3}')
            previous = default.read_bytes()
            model = fixture()
            with patch.object(step1, "ROOT", root), patch.object(
                    step1, "prepare_models", return_value={"Case": (model, pfba(model), ["A[e]"])}):
                step1.main(["--plan", "results/results_by_AI/case/analysis_plan.json"])
            plan = json.loads((root / "results/results_by_AI/case/analysis_plan.json").read_text())
            self.assertEqual(plan["output_dir"], "results/results_by_AI/case")
            self.assertTrue((root / plan["models"]["Case"]["input_file"]).is_file())
            self.assertEqual(default.read_bytes(), previous)
    def test_copied_example_uses_flat_case_directory(self):
        import importlib.util
        import shutil
        import sys

        with tempfile.TemporaryDirectory() as directory, patch("sys.path", list(sys.path)):
            root = Path(directory)
            marker = root / "carbon_flow_tools/model_preparation.py"
            marker.parent.mkdir()
            marker.touch()
            script = root / "results/results_by_AI/custom/step1_prepare_inputs.py"
            script.parent.mkdir(parents=True)
            template = Path(__file__).resolve().parents[1] / "examples/step1_prepare_custom_model_inputs.py"
            shutil.copy2(template, script)
            spec = importlib.util.spec_from_file_location("custom_case_example", script)
            example = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(example)
            self.assertEqual(example.ROOT, root)
            self.assertEqual(example.CASE_DIR, script.parent)
            model = fixture()
            with patch.object(example, "prepare_models", return_value={
                "Custom": (model, pfba(model), ["A[e]"])
            }), patch("carbon_flow_tools.cli.ROOT", root):
                example.main([])
                for stage, defaults in (("map", {}), ("trace", {"rounds": 1})):
                    settings, jobs = load_analysis_plan(script.parent / "analysis_plan.json", stage, defaults)
                    self.assertEqual(settings["output_dir"], "results/results_by_AI/custom")
                    self.assertIn("Custom", jobs)
            plan = json.loads((script.parent / "analysis_plan.json").read_text())
            self.assertEqual((root / plan["models"]["Custom"]["input_file"]).parent, script.parent)
            self.assertFalse((script.parent / "results").exists())

if __name__ == "__main__":
    unittest.main()
