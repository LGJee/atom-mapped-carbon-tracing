"""Regression checks for preserved scientific behavior."""

import hashlib
import os
import time
from unittest.mock import patch
import sys
import tempfile
from zipfile import ZipFile
from copy import deepcopy
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from rdkit import Chem

from carbon_flow_tools.config import parallel_capacity
from carbon_flow_tools.consensus import mapping_signatures, select_mappings
from carbon_flow_tools.pipeline import run_analyses, read_input, reuse_predictions
from carbon_flow_tools.reactions import (
    PreparedReaction,
    build_side,
    parse_side,
    prepare_reactions,
)


def load_model(options):
    events = Path(options["events"])
    (events / (str(os.getpid()) + ".started")).touch()
    deadline = time.monotonic() + 15
    while len(list(events.glob("*.started"))) < 6:
        if time.monotonic() > deadline:
            raise RuntimeError(
                "Six cross-strain predictions did not start concurrently"
            )
        time.sleep(0.01)
    return options["mapping"], options["score"], options


def map_batch(model, smiles):
    options = model[2]
    options["processed"] = options.get("processed", 0) + len(smiles)
    if options["processed"] == len(options["reactions"]):
        (Path(options["events"]) / (str(os.getpid()) + ".done")).touch()
    return [{"mapped_rxn": model[0], "confidence": model[1]} for _ in smiles]


def reaction(mapped, smiles="CCO", left="A[c]", right="B[c]", score=True):
    structures = {left: smiles, right: smiles}
    l, error = build_side(parse_side(left), structures)
    assert not error
    r, error = build_side(parse_side(right), structures)
    assert not error
    return PreparedReaction(
        "r_test",
        "",
        1,
        left + " -> " + right,
        smiles + ">>" + smiles,
        l,
        r,
        source_text=left,
        status="mapped",
        mapped_smiles=mapped,
        confident=score,
    )


def main():
    with patch(
        "carbon_flow_tools.config.os.sched_getaffinity", return_value=set(range(32))
    ), patch(
        "carbon_flow_tools.config.Path.read_text",
        return_value="MemAvailable: 33554432 kB\n",
    ):
        assert parallel_capacity(6) == (4, 8)
        assert parallel_capacity(6, memory_per_process_gb=4) == (6, 5)
        assert parallel_capacity(6, max_processes=2, memory_per_process_gb=4) == (2, 16)
        assert parallel_capacity(1) == (1, 32)
    with patch(
        "carbon_flow_tools.config.os.sched_getaffinity", return_value={0, 1}
    ), patch(
        "carbon_flow_tools.config.Path.read_text",
        return_value="MemAvailable: 33554432 kB\n",
    ):
        assert parallel_capacity(6, memory_per_process_gb=4) == (2, 1)
    with patch("carbon_flow_tools.config.Path.read_text", side_effect=OSError):
        assert parallel_capacity(6)[0] == 1
    for kwargs in (
        {"max_processes": 0},
        {"memory_per_process_gb": 0},
        {"memory_per_process_gb": float("nan")},
    ):
        try:
            parallel_capacity(6, **kwargs)
        except ValueError:
            pass
        else:
            raise AssertionError("invalid resource configuration accepted")
    correct = "[CH3:1][CH2:2][OH:3]>>[CH3:1][CH2:2][OH:3]"
    alias = "[OH:90][CH2:80][CH3:70]>>[CH3:70][CH2:80][OH:90]"
    swapped = "[CH3:1][CH2:2][OH:3]>>[CH3:2][CH2:1][OH:3]"
    assert mapping_signatures(reaction(correct)) == mapping_signatures(reaction(alias))
    assert (
        mapping_signatures(reaction(correct))[0]
        != mapping_signatures(reaction(swapped))[0]
    )

    assert mapping_signatures(
        reaction("[CH3:1][CH3:2]>>[CH3:2][CH3:1]", "CC")
    ) == mapping_signatures(reaction("[CH3:4][CH3:5]>>[CH3:4][CH3:5]", "CC"))
    duplicate = reaction(correct + "." + correct)
    for side in (duplicate.left, duplicate.right):
        side.extend(deepcopy(side))
    duplicate.mapped_smiles = "[CH3:1][CH2:2][OH:3].[CH3:4][CH2:5][OH:6]>>[CH3:1][CH2:2][OH:3].[CH3:4][CH2:5][OH:6]"
    other = deepcopy(duplicate)
    other.mapped_smiles = (
        duplicate.mapped_smiles.split(">>")[0]
        + ">>[CH3:4][CH2:5][OH:6].[CH3:1][CH2:2][OH:3]"
    )
    assert mapping_signatures(duplicate) == mapping_signatures(other)

    for raw in (
        "[CH3:1]CO>>[CH3:1]CO",
        "[CH3:1][CH2:2][OH:3]>>[CH3:1][CH2:4][OH:3]",
        "[CH3:1][CH2:2][OH:3]>>[CH3:1][CH2:3][OH:2]",
    ):
        try:
            mapping_signatures(reaction(raw))
        except ValueError:
            pass
        else:
            raise AssertionError("invalid carbon mapping accepted: " + raw)
    candidates = {
        "LocalMapper": [reaction(correct)],
        "RXNMapper": [reaction(alias, score=0.99)],
        "EnzymeMap": [reaction(swapped, score=0.98)],
    }
    selected, table = select_mappings(candidates)
    assert table.iloc[0].SelectedMapper == "LocalMapper"
    assert len(table.iloc[0].SupportingMappers.split("; ")) == 2 and not bool(
        table.iloc[0].ReviewRequired
    )
    assert selected[0].mapped_smiles == correct
    candidates.pop("RXNMapper")
    selected, table = select_mappings(candidates)
    assert table.iloc[0].SelectedMapper == "EnzymeMap" and bool(
        table.iloc[0].ReviewRequired
    )
    assert selected[0].mapped_smiles == swapped
    candidates = {
        "LocalMapper": [reaction(correct)],
        "RXNMapper": [reaction(swapped, score=0.99)],
    }
    assert select_mappings(candidates)[1].iloc[0].SelectedMapper == "LocalMapper"
    candidates["LocalMapper"][0].confident = False
    assert select_mappings(candidates)[1].iloc[0].SelectedMapper == "RXNMapper"
    for item in candidates.values():
        item[0].status = "failed"
    selected, table = select_mappings(candidates)
    assert table.iloc[0].SelectedMapper == "" and not selected[0].mapped_smiles
    assert bool(table.iloc[0].ReviewRequired)
    assert (
        table.iloc[0].SelectionReason == "No mapping available for carbon propagation"
    )

    original = Chem.MolFromSmiles("C[C@H](O)CCN")
    for atom in original.GetAtoms():
        atom.SetAtomMapNum(atom.GetIdx() + 1)
    a = Chem.MolToSmiles(original, canonical=False, rootedAtAtom=0)
    b = Chem.MolToSmiles(original, canonical=False, rootedAtAtom=5)
    assert mapping_signatures(
        reaction(a + ">>" + a, "C[C@H](O)CCN")
    ) == mapping_signatures(reaction(b + ">>" + a, "C[C@H](O)CCN"))

    a = "[CH3:1][C:2](=[O:3])[OH:4]>>[CH3:1][C:2](=[O:3])[OH:4]"
    b = "[CH3:1][C:2](=[O:3])[OH:4]>>[CH3:1][C:2](=[O:4])[OH:3]"
    selected, table = select_mappings(
        {"EnzymeMap": [reaction(a, "CC(=O)O")], "RXNMapper": [reaction(b, "CC(=O)O")]}
    )
    assert (
        len(table.iloc[0].SupportingMappers.split("; ")) == 2
        and table.iloc[0].HeavyAtomAgreement == "disagreement"
    )
    assert not bool(table.iloc[0].ReviewRequired)
    assert (
        table.iloc[0].SelectionReason
        == "2 tools agree on carbon mapping; selected directly"
    )
    _, single = select_mappings({"RXNMapper": [reaction(correct, score=0.99)]})
    assert bool(single.iloc[0].ReviewRequired)
    with tempfile.TemporaryDirectory() as folder:
        root = Path(folder)
        source = root / "input.xlsx"
        wb = Workbook()
        wb.active.title = "non_IUPD_nonzero"
        wb.active.append(["ReactionID", "ReactionName", "Equation", "Flux"])
        wb.active.append(["r1", "transport", "A[e] -> A[c]", 1])
        wb.active.append(["r2", "reaction", "A[c] -> B[c]", 1])
        met = wb.create_sheet("Metabolites")
        met.append(["MetaboliteID", "MetaboliteName", "SMILES_Yeast9", "SMILES_Manual"])
        for i, name in enumerate(("A[e]", "A[c]", "B[c]")):
            met.append(["s_" + str(i), name, "CCO", ""])
        wb.save(source)

        precise = source.with_suffix(".precision.xlsx")
        with ZipFile(source) as old, ZipFile(precise, "w") as new:
            for item in old.infolist():
                data = old.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    needle = b'<c r="D2" t="n"><v>1</v>'
                    assert needle in data
                    data = data.replace(
                        needle, b'<c r="D2" t="n"><v>1.3979095809545714</v>'
                    )
                    data = data.replace(
                        b'<c r="D3" t="n"><v>1</v>', b'<c r="D3" t="n"><v>1e-16</v>'
                    )
                new.writestr(item, data)
        precise.replace(source)
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        names = ["LocalMapper", "RXNMapper", "EnzymeMap"]
        events = root / "events"
        events.mkdir()
        settings = {
            name: {
                "module": "scripts.test_consensus_mapping",
                "python": sys.executable,
                "options": {
                    "mapping": correct,
                    "score": True if name == "LocalMapper" else 0.8,
                    "events": str(events),
                },
            }
            for name in names
        }
        job = {
            "input_file": source,
            "reaction_sheet": "non_IUPD_nonzero",
            "metabolite_sheet": "Metabolites",
            "carbon_sources": ["A[e]"],
        }

        def checked_selection(predictions):
            assert len(list(events.glob("*.done"))) == 6, (
                "Selection started before all six predictions completed"
            )
            return select_mappings(predictions)

        with patch(
            "carbon_flow_tools.pipeline.select_mappings", side_effect=checked_selection
        ), patch("carbon_flow_tools.pipeline.parallel_capacity", return_value=(6, 1)):
            for _ in range(2):
                run_analyses(
                    {"non_IUPD": job, "IUPD": dict(job)},
                    2,
                    1,
                    800,
                    root / "out",
                    names,
                    settings,
                )
        assert len(list(events.glob("*.started"))) == 6, (
            "Incorrect worker isolation or prediction reuse"
        )
        iupd_output = root / "out" / "IUPD_2_atom_mapping.xlsx"
        with pd.ExcelFile(iupd_output, engine="openpyxl") as excel:
            assert len(pd.read_excel(excel, sheet_name="Final_Mapping")) == 2
            assert (
                len(
                    pd.read_excel(
                        root / "out/non_IUPD_3_carbon_flow.xlsx",
                        sheet_name="Carbon_Flow_R1",
                        engine="openpyxl",
                    )
                )
                == 2
            )
        output = root / "out" / "non_IUPD_2_atom_mapping.xlsx"
        with pd.ExcelFile(output, engine="openpyxl") as excel:
            assert set(names + ["Mapping_Selection", "Final_Mapping"]) <= set(
                excel.sheet_names
            )
            for name in names:
                assert pd.read_excel(excel, sheet_name=name).columns.tolist() == [
                    "ReactionID",
                    "ReactionName",
                    "Equation",
                    "ReactionSMILES",
                    "MappedReactionSMILES",
                    "MappingDirection",
                    "MappingConfident",
                ]
            info = pd.read_excel(excel, sheet_name="Mapping_Info")
            assert info.ReusedPredictions.all()
            choices = pd.read_excel(excel, sheet_name="Mapping_Selection")
            final = pd.read_excel(excel, sheet_name="Final_Mapping")
            flow = pd.read_excel(
                root / "out/non_IUPD_3_carbon_flow.xlsx",
                sheet_name="Carbon_Flow_R1",
                engine="openpyxl",
            )
            assert choices.SelectedMapper.tolist() == ["Identity", "EnzymeMap"]
            assert "CarbonAgreement" not in choices and not choices.ReviewRequired.any()
            assert len(final) == 2 and final.MappedReactionSMILES.eq(correct).all()
            assert set(flow.Production) == {"A[c]", "B[c]"}
            assert flow.CarbonLabels.eq("C1,C2").all()
        original_output = output.read_bytes()
        completed = []

        def failed_mapper(strain, reactions, name, settings, cpu_threads=1):
            if name == "EnzymeMap":
                raise RuntimeError("Test worker failed")
            time.sleep(0.05)
            completed.append(name)

        iupd_before = iupd_output.read_bytes()
        with patch(
            "carbon_flow_tools.pipeline.run_mapper", side_effect=failed_mapper
        ), patch("carbon_flow_tools.pipeline.parallel_capacity", return_value=(6, 1)):
            try:
                run_analyses(
                    {"non_IUPD": job, "IUPD": dict(job)},
                    1,
                    1,
                    800,
                    root / "out",
                    names,
                    settings,
                    refresh=True,
                )
            except RuntimeError as error:
                assert str(error) == "Test worker failed"
            else:
                raise AssertionError("Selection continued after worker failure")
        assert set(completed) == {"LocalMapper", "RXNMapper"}
        assert output.read_bytes() == original_output
        assert iupd_output.read_bytes() == iupd_before
        assert hashlib.sha256(source.read_bytes()).hexdigest() == before
        cached = load_workbook(output)
        headers = {cell.value: cell.column for cell in cached["LocalMapper"][1]}
        cached["LocalMapper"].cell(2, headers["MappedReactionSMILES"], swapped)
        cached.save(output)
        cached.close()
        transports, original_table, original_metabolites = read_input(
            source, "non_IUPD_nonzero", "Metabolites"
        )
        assert reuse_predictions(
            output,
            "LocalMapper",
            settings["LocalMapper"],
            transports,
            original_table,
            original_metabolites,
            job,
        )
        assert transports[0].is_transport and transports[0].mapped_smiles == correct
        _, source_table, metabolites = read_input(
            source, "non_IUPD_nonzero", "Metabolites"
        )
        source_table.loc[1, "Flux"] = -1e-16
        reverse = prepare_reactions(source_table, metabolites)
        assert not reuse_predictions(
            output,
            "LocalMapper",
            settings["LocalMapper"],
            reverse,
            source_table,
            metabolites,
            job,
        )
        protected = dict(job, input_file=output)
        try:
            run_analyses(
                {"non_IUPD": protected}, 1, 1, 800, root / "out", names, settings
            )
        except ValueError as error:
            assert "input" in str(error)
        else:
            raise AssertionError("input overwrite allowed")
    print(
        "PASS: mapping equivalence, complete carbon coverage, consensus, score policy, reuse, parallel processes, completion barrier, failure protection, direction and carbon flow"
    )


if __name__ == "__main__":
    main()
