"""Check Summary copies, safe sheet names and replacement of the aggregate."""

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from carbon_flow_tools.pipeline import collect_carbon_summaries


class SummaryCollectionTests(unittest.TestCase):
    def test_copies_all_saved_variants_and_rebuilds(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            names = ["Long_model_" + "x" * 40 + ending for ending in ("a", "b")]
            names += [names[0] + "_AIGenerated", "short[model]"]
            expected = []
            hashes = {}
            for index, name in enumerate(names):
                model = name.replace("_AIGenerated", "")
                suffix = "_AIGenerated" if name.endswith("_AIGenerated") else ""
                path = root / (model + "_3_carbon_flow" + suffix + ".xlsx")
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "Summary"
                rows = [("Round", "A_LabeledCarbonMetabolites", "StopReason"),
                        (1, index + 2, None), (2, index + 4, "round_limit")]
                for row in rows:
                    sheet.append(row)
                workbook.create_sheet("Carbon_Flow").append(["not copied"])
                workbook.save(path)
                workbook.close()
                hashes[path] = path.read_bytes()
                expected.append(rows)
            # Unrelated, lock and unfinished candidate workbooks must be ignored.
            for name in ("unrelated.xlsx", "old_3_carbon_flow_Enumeration.xlsx",
                         "unfinished_3_carbon_flow.tmp.xlsx",
                         "~$locked_3_carbon_flow.xlsx"):
                (root / name).write_bytes(b"ignored")
            nested = root / "other_case"
            nested.mkdir()
            (nested / "nested_3_carbon_flow.xlsx").write_bytes(b"ignored")
            collect_carbon_summaries(root)
            output = root / "carbon_flow_summary.xlsx"
            for _ in range(2):
                workbook = load_workbook(output)
                self.assertEqual(len(workbook.worksheets), len(expected))
                self.assertEqual(len({s.title.lower() for s in workbook}), len(expected))
                self.assertTrue(any(s.title.endswith("_AI") for s in workbook))
                actual = [list(s.values) for s in workbook]
                self.assertCountEqual(actual, expected)
                for sheet in workbook:
                    self.assertLessEqual(len(sheet.title), 31)
                    self.assertEqual(sheet["A1"].font.name, "Calibri")
                    self.assertIsNone(sheet["A1"].font.scheme)
                workbook.close()
                collect_carbon_summaries(root)
            for path, before in hashes.items():
                self.assertEqual(path.read_bytes(), before)
            before = output.read_bytes()
            broken = root / "broken_3_carbon_flow.xlsx"
            workbook = Workbook()
            workbook.save(broken)
            workbook.close()
            with self.assertRaises(KeyError):
                collect_carbon_summaries(root)
            self.assertEqual(output.read_bytes(), before)
            self.assertFalse(output.with_suffix(".tmp.xlsx").exists())

    def test_no_results_creates_no_empty_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            collect_carbon_summaries(root)
            self.assertEqual(list(root.iterdir()), [])


if __name__ == "__main__":
    unittest.main()
