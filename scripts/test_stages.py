"""Check mapping-only output and propagation without any mapper calls."""

import hashlib
from pathlib import Path
import sys
import tempfile
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from carbon_flow_tools.pipeline import (
    run_analyses,
    run_carbon_flow,
    read_input,
    reuse_predictions,
)
from scripts.test_parallel_carbon import make_input


def main():
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        source = root / "input.xlsx"
        make_input(source, "B[c]")
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        jobs = {
            "Example": dict(
                input_file=source,
                reaction_sheet="Nonzero",
                metabolite_sheet="Metabolites",
                carbon_sources=["A[e]"],
            )
        }
        configs = {
            "ExampleMapper": {
                "module": "scripts.test_parallel_carbon",
                "python": sys.executable,
                "options": {"test": True},
            }
        }
        arguments = (jobs, 2, 1, 800)
        options = dict(max_mapper_processes=1, max_carbon_processes=1)
        try:
            run_analyses(
                *arguments,
                root / "missing",
                ["ExampleMapper"],
                {},
                stage="trace",
                **options,
            )
        except FileNotFoundError as error:
            assert "step2_atom_mapping.py" in str(error)
        else:
            raise AssertionError("Propagation accepted a missing step 2 workbook")
        run_analyses(
            *arguments, root / "combined", ["ExampleMapper"], configs, **options
        )
        run_analyses(
            *arguments,
            root / "split",
            ["ExampleMapper"],
            configs,
            stage="map",
            **options,
        )
        mapping_output = root / "split/Example_2_atom_mapping.xlsx"
        output = root / "split/Example_3_carbon_flow.xlsx"
        mapping_before = mapping_output.read_bytes()
        with pd.ExcelFile(mapping_output, engine="openpyxl") as excel:
            assert not output.exists()
            assert "Final_Mapping" in excel.sheet_names
            assert not any(name.startswith("Carbon_Flow") for name in excel.sheet_names)
            selection = pd.read_excel(excel, sheet_name="Mapping_Selection")
            assert "AIGenerated" not in selection.columns
            final = pd.read_excel(excel, sheet_name="Final_Mapping")
            assert (
                final.columns[final.columns.get_loc("ReviewRequired") + 1]
                == "AIGenerated"
            )
            assert final.AIGenerated.eq(False).all()
        with patch(
            "carbon_flow_tools.pipeline.run_mapper",
            side_effect=AssertionError("Mapper called during trace"),
        ):
            run_analyses(
                *arguments,
                root / "split",
                ["ExampleMapper"],
                {},
                stage="trace",
                **options,
            )
        assert mapping_output.read_bytes() == mapping_before
        with pd.ExcelFile(output, engine="openpyxl") as split, pd.ExcelFile(
            root / "combined/Example_3_carbon_flow.xlsx", engine="openpyxl"
        ) as combined:
            assert set(split.sheet_names) == {
                "Carbon_Flow_R1",
                "Carbon_Flow_R2",
                "Summary",
            }
            for name in (
                "Carbon_Flow_R1",
                "Carbon_Flow_R2",
                "Summary",
            ):
                pd.testing.assert_frame_equal(
                    pd.read_excel(split, sheet_name=name),
                    pd.read_excel(combined, sheet_name=name),
                )
        assert hashlib.sha256(source.read_bytes()).hexdigest() == before
        # Step 3 is also usable when step 1 and mapper environments are unavailable.
        independent_jobs = {
            "Example": {"mapping_file": mapping_output, "carbon_sources": ["A[e]"]}
        }
        hidden_source = source.with_suffix(".hidden.xlsx")
        source.rename(hidden_source)
        try:
            with patch(
                "carbon_flow_tools.pipeline.run_mapper",
                side_effect=AssertionError("Mapper called during independent trace"),
            ):
                run_carbon_flow(
                    independent_jobs, 2, root / "independent", max_processes=1
                )
            with pd.ExcelFile(
                root / "independent/Example_3_carbon_flow.xlsx", engine="openpyxl"
            ) as independent:
                assert set(independent.sheet_names) == {
                    "Carbon_Flow_R1",
                    "Carbon_Flow_R2",
                    "Summary",
                }
                for sheet in independent.sheet_names:
                    pd.testing.assert_frame_equal(
                        pd.read_excel(independent, sheet_name=sheet),
                        pd.read_excel(output, sheet_name=sheet, engine="openpyxl"),
                    )
            parallel_jobs = {
                name: dict(independent_jobs["Example"]) for name in ("One", "Two")
            }
            run_carbon_flow(
                parallel_jobs,
                2,
                root / "parallel_trace",
                max_processes=2,
                memory_gb=0.25,
            )
            for name in parallel_jobs:
                with pd.ExcelFile(
                    root / "parallel_trace" / (name + "_3_carbon_flow.xlsx"),
                    engine="openpyxl",
                ) as parallel:
                    for sheet in parallel.sheet_names:
                        pd.testing.assert_frame_equal(
                            pd.read_excel(parallel, sheet_name=sheet),
                            pd.read_excel(output, sheet_name=sheet, engine="openpyxl"),
                        )
            assert mapping_output.read_bytes() == mapping_before
        finally:
            hidden_source.rename(source)
        run_analyses(
            jobs,
            1,
            1,
            800,
            root / "split",
            ["ExampleMapper"],
            {},
            stage="trace",
            **options,
        )
        with pd.ExcelFile(output, engine="openpyxl") as excel:
            assert "Carbon_Flow_R2" not in excel.sheet_names
        # Numeric-looking model IDs must survive the Excel mapping/trace boundary.
        numeric_mapping = root / "numeric_mapping.xlsx"
        workbook = load_workbook(mapping_output)
        for cells in workbook:
            columns = [cell.column for cell in cells[1] if cell.value == "ReactionID"]
            if not columns:
                continue
            for row in range(2, cells.max_row + 1):
                cell = cells.cell(row, columns[0])
                if (
                    isinstance(cell.value, str)
                    and cell.value.startswith("r")
                    and cell.value[1:].isdigit()
                ):
                    cell.value = int(cell.value[1:])
        workbook.save(numeric_mapping)
        workbook.close()
        numeric_reactions, numeric_source, numeric_metabolites = read_input(
            numeric_mapping, "Nonzero", "Metabolites"
        )
        assert reuse_predictions(
            numeric_mapping,
            "ExampleMapper",
            configs["ExampleMapper"],
            numeric_reactions,
            numeric_source,
            numeric_metabolites,
            jobs["Example"],
        )
        run_carbon_flow(
            {"Numeric": dict(mapping_file=numeric_mapping, carbon_sources=["A[e]"])},
            1,
            root / "numeric",
            max_processes=1,
        )
        numeric_flow = pd.read_excel(
            root / "numeric/Numeric_3_carbon_flow.xlsx",
            sheet_name="Carbon_Flow_R1",
            engine="openpyxl",
        )
        assert len(numeric_flow) == 2
        assert set(numeric_flow.ReactionID.astype(str)) == {"1", "2"}

        for kind in (
            "duplicate_id",
            "zero_flux",
            "nonfinite_flux",
            "conflicting_structure",
        ):
            invalid = root / (kind + ".xlsx")
            workbook = load_workbook(source)
            if kind == "duplicate_id":
                workbook["Nonzero"]["A3"] = "r1"
            elif kind in ("zero_flux", "nonfinite_flux"):
                workbook["Nonzero"]["D2"] = 0 if kind == "zero_flux" else float("nan")
            else:
                workbook["Metabolites"]["B4"] = "A[c]"
                workbook["Metabolites"]["C4"] = "CC"
            workbook.save(invalid)
            workbook.close()
            try:
                read_input(invalid, "Nonzero", "Metabolites")
            except ValueError:
                pass
            else:
                raise AssertionError("Invalid input accepted: " + kind)
    print(
        "PASS: separate stages, no prediction during trace, flags, fewer-round cleanup, input preservation"
    )


def check_mapping_variants():
    import json
    import shutil
    from carbon_flow_tools.cli import main, discover_inputs
    from carbon_flow_tools.atom_helpers import assignment_from_labeled_smiles
    import step3_carbon_flow as tracing

    with tempfile.TemporaryDirectory() as directory:
        folder = Path(directory)
        source = folder / "input.xlsx"
        make_input(source, "B[c]")
        source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        names = ["Mixed_Culture_non_IUPD", "Mixed_Culture_IUPD"]
        jobs = {
            name: dict(
                input_file=source,
                input_sha256=source_hash,
                reaction_sheet="Nonzero",
                metabolite_sheet="Metabolites",
            )
            for name in names
        }
        run_analyses(
            jobs,
            1,
            1,
            800,
            folder,
            ["ExampleMapper"],
            {
                "ExampleMapper": {
                    "module": "scripts.test_parallel_carbon",
                    "python": sys.executable,
                    "options": {"test": True},
                }
            },
            stage="map",
            max_mapper_processes=1,
        )
        plan = folder / "analysis_plan.json"
        plan.write_text(
            json.dumps(
                dict(
                    version=3,
                    output_dir=str(folder),
                    models={
                        name: dict(
                            input_file=str(source),
                            input_sha256=source_hash,
                            carbon_sources=["A[e]"],
                        )
                        for name in names
                    },
                )
            ),
            encoding="utf-8",
        )
        options = tracing.settings()
        options.update(
            analysis_plan=str(plan),
            tracing_python=sys.executable,
            rounds=2,
            max_carbon_processes=1,
        )
        assert main("trace", options, [names[0]]) == 0
        original_flow = folder / (names[0] + "_3_carbon_flow.xlsx")
        originals = {
            path: hashlib.sha256(path.read_bytes()).hexdigest()
            for path in folder.glob("*.xlsx")
        }
        reviewed = folder / (names[0] + "_2_atom_mapping_AIGenerated.xlsx")
        shutil.copyfile(folder / (names[0] + "_2_atom_mapping.xlsx"), reviewed)
        workbook = load_workbook(reviewed)
        sheet = workbook["Final_Mapping"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, headers["ReactionID"]).value == "r2":
                sheet.cell(
                    row,
                    headers["MappedReactionSMILES"],
                    "[CH3:1][CH2:2][OH:3]>>[CH3:2][CH2:1][OH:3]",
                )
                sheet.cell(row, headers["AIGenerated"], True)
        workbook.save(reviewed)
        workbook.close()
        reviewed_hash = hashlib.sha256(reviewed.read_bytes()).hexdigest()
        assert set(discover_inputs(folder, "trace")) == set(names)
        with patch(
            "carbon_flow_tools.pipeline.run_mapper",
            side_effect=AssertionError("Step 3 called a mapper"),
        ):
            assert main("trace", options, []) == 0
        report = json.loads((folder / "carbon_flow_run.json").read_text())
        assert len(report["models"]) == 3
        assert sum(item["model"] == names[0] for item in report["models"]) == 2
        assert len({item["output_file"] for item in report["models"]}) == 3
        assert {p.name for p in folder.glob("*_3_carbon_flow*.xlsx")} == {
            names[0] + "_3_carbon_flow.xlsx",
            names[0] + "_3_carbon_flow_AIGenerated.xlsx",
            names[1] + "_3_carbon_flow.xlsx",
        }
        assert all(item["mapping_sha256"] == hashlib.sha256(
            Path(item["mapping_file"]).read_bytes()).hexdigest() for item in report["models"])
        aggregate = folder / "carbon_flow_summary.xlsx"
        with pd.ExcelFile(aggregate, engine="openpyxl") as summaries:
            assert len(summaries.sheet_names) == 3
            for result in report["models"]:
                path = Path(result["output_file"])
                title = result["model"] + ("_AI" if path.stem.endswith("_AIGenerated") else "")
                pd.testing.assert_frame_equal(
                    pd.read_excel(summaries, sheet_name=title),
                    pd.read_excel(path, sheet_name="Summary", engine="openpyxl"),
                )
        updated_flow = folder / (names[0] + "_3_carbon_flow_AIGenerated.xlsx")
        table = pd.read_excel(
            updated_flow, sheet_name="Carbon_Flow_R1", engine="openpyxl"
        )
        marked = table.loc[table.Production == "B[c]", "LabeledSMILES"].iloc[0]
        assert assignment_from_labeled_smiles("CCO", marked) == (
            (0, 0, "C2"),
            (0, 1, "C1"),
        )
        before = pd.read_excel(
            original_flow, sheet_name="Carbon_Flow_R1", engine="openpyxl"
        )
        assert (
            before.loc[before.Production == "B[c]", "LabeledSMILES"].iloc[0] != marked
        )
        assert pd.ExcelFile(updated_flow, engine="openpyxl").sheet_names == [
            "Summary",
            "Carbon_Flow_R1",
            "Carbon_Flow_R2",
            "EMU_Metabolites",
            "EMU_BoundaryInputs",
        ]
        assert hashlib.sha256(reviewed.read_bytes()).hexdigest() == reviewed_hash
        for path, digest in originals.items():
            if path.name != "carbon_flow_summary.xlsx" and "_3_carbon_flow" not in path.stem:
                assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
        # The fallback interpreter must relaunch the unified propagation entry.
        with patch("carbon_flow_tools.cli.sys.executable", "/unrelated/python"), patch(
            "carbon_flow_tools.cli.subprocess.run"
        ) as relaunch:
            relaunch.return_value.returncode = 0
            assert main("trace", options, [names[0], "--rounds", "1"]) == 0
            assert Path(relaunch.call_args[0][0][1]).name == "step3_carbon_flow.py"
        second = folder / (names[1] + "_2_atom_mapping_AIGenerated.xlsx")
        shutil.copyfile(reviewed, second)
        assert set(discover_inputs(folder, "trace")) == set(names)
        assert main("trace", options, []) == 0
        assert (folder / (names[1] + "_3_carbon_flow_AIGenerated.xlsx")).exists()
        # Parallel tasks for two versions of the same model remain isolated.
        parallel = folder / "parallel"
        run_carbon_flow({names[0]: dict(
            mapping_files=[folder / (names[0] + "_2_atom_mapping.xlsx"), reviewed],
            carbon_sources=["A[e]"],
        )}, 2, parallel, max_processes=2, memory_gb=.25)
        for filename in (names[0] + "_3_carbon_flow.xlsx", updated_flow.name):
            for sheet_name in ("Summary", "Carbon_Flow_R1", "Carbon_Flow_R2"):
                pd.testing.assert_frame_equal(
                    pd.read_excel(folder / filename, sheet_name=sheet_name, engine="openpyxl"),
                    pd.read_excel(parallel / filename, sheet_name=sheet_name, engine="openpyxl"),
                )
        # A model can have only an AI workbook; normal output is not overwritten.
        normal = folder / (names[1] + "_2_atom_mapping.xlsx")
        saved_input = normal.read_bytes()
        normal.unlink()
        with patch("carbon_flow_tools.pipeline.run_carbon_flow") as propagate:
            assert main("trace", options, [names[1]]) == 0
            assert propagate.call_args[0][0][names[1]]["mapping_files"] == [second]
        normal.write_bytes(saved_input)
        # A stale AI copy must fail before replacing any previous good result.
        output_hash = hashlib.sha256(updated_flow.read_bytes()).hexdigest()
        workbook = load_workbook(reviewed)
        sheet = workbook["Mapping_Info"]
        column = [c.value for c in sheet[1]].index("SourceInputSHA256") + 1
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, column).value:
                sheet.cell(row, column, "0" * 64)
        workbook.save(reviewed)
        workbook.close()
        try:
            main("trace", options, [names[0]])
        except ValueError as error:
            assert "does not match" in str(error)
        else:
            raise AssertionError("Stale AI mapping was accepted")
        assert hashlib.sha256(updated_flow.read_bytes()).hexdigest() == output_hash
        reviewed.unlink()
        second.unlink()
        # Normal-only inputs work; previous AI results remain available.
        assert main("trace", options, []) == 0
        assert hashlib.sha256(updated_flow.read_bytes()).hexdigest() == output_hash
    print(
        "PASS: AI copies, changed label transfer, separate outputs, hashes and relaunch"
    )


def check_calibri_export():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from carbon_flow_tools.pipeline import save_output

    with tempfile.TemporaryDirectory() as directory:
        output = Path(directory) / "font.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["ReactionID", "Equation"])
        sheet.append(["r1", "A[c] -> B[c]"])
        sheet["A1"].font = Font(name="SimSun", scheme="minor", charset=134, bold=True)
        sheet["B2"].alignment = Alignment(wrap_text=True)
        save_output(workbook, output)
        # Adding a font after a save must still use a consistent indexed table.
        sheet["A2"].font = Font(name="Arial", italic=True)
        save_output(workbook, output)
        workbook.close()
        workbook = load_workbook(output)
        for font in workbook._fonts:
            assert (
                font.name == "Calibri" and font.scheme is None and font.charset is None
            )
        assert workbook.active["A1"].font.bold
        assert workbook.active["A2"].font.italic
        assert workbook.active["B2"].alignment.wrap_text
        assert workbook.active["B2"].value == "A[c] -> B[c]"
        workbook.close()
    print(
        "PASS: explicit Calibri export, repeated saves and unchanged cell values/styles"
    )


if __name__ == "__main__":
    main()
    check_mapping_variants()
    check_calibri_export()
