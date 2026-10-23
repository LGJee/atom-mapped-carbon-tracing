"""Check bounded states, unchanged snapshots, compact rounds and safe stops."""
from contextlib import redirect_stdout
from io import StringIO
from itertools import islice, product
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
import pandas as pd

from carbon_flow_tools.carbon_trace import propagate_rounds, projected_choices, map_one_reaction
from carbon_flow_tools.config import DiskGuard, DiskSpaceLimitError, MemoryPressureError
from carbon_flow_tools.flow_output import FlowOutput
from carbon_flow_tools.pipeline import trace_analysis, summarize_sources
from carbon_flow_tools.propagation_state import ordered_product, PAGE_SIZE, CACHE_PAGES, CATALOG_CACHE
from scripts.test_carbon_tracing import reaction


def fixture():
    structures = {'A[c]': 'CO', 'B[c]': 'C', 'P[c]': 'CCO'}
    mapped = '[CH3:1][OH:3].[CH4:2]>>[CH3:2][CH2:1][OH:3]'
    reactions = []
    for index, (left, right, text) in enumerate([
        ('A[c] + B[c]', 'P[c]', mapped),
        ('P[c]', 'A[c] + B[c]', '>>'.join(reversed(mapped.split('>>')))),
        ('P[c]', 'P[c]', '[CH3:2][CH2:1][OH:3]>>[CH3:1][CH2:2][OH:3]')]):
        rxn, _ = reaction(left, right, structures, text)
        rxn.status, rxn.reaction_id = 'mapped', 'r' + str(index)
        reactions.append(rxn)
    metabolites = pd.DataFrame([dict(MetaboliteID=k, MetaboliteName=k, SMILES_Model=v)
                               for k, v in structures.items()])
    return structures, reactions, metabolites


class CompactRoundTests(unittest.TestCase):
    def make_store(self, folder, detail=3, batch=2):
        return FlowOutput(Path(folder)/'out.xlsx', 50_000_000, batch, [], None, detail_rounds=detail)

    def test_ordered_product_matches_product_without_consuming_pools(self):
        for choices in ([], [[1, 2], [None, 3], [4, 5]], [[], [1]], [[1]]):
            self.assertEqual(list(ordered_product(choices)), list(product(*choices)))
        visits = []
        class LargeChoices:
            def __iter__(self):
                for value in range(1000000):
                    visits.append(value)
                    yield value
        self.assertEqual(list(islice(ordered_product([LargeChoices(), LargeChoices()]), 3)),
                         [(0, 0), (0, 1), (0, 2)])
        self.assertEqual(len(visits), 4)

    def test_paged_states_match_exact_projection_and_invalidate_earlier_steps(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.make_store(folder)
            try:
                states, steps = {}, {}
                for number in range(PAGE_SIZE * (CACHE_PAGES + 2)):
                    assignment = ((0, 0, 'C' + str(number % 7 + 1)),)
                    states[number], steps[('A', number)] = assignment, number % 5
                    store.states['A'][number] = assignment
                    store.first_steps[('A', number)] = number % 5
                groups = [(0, 'A', ((0, 0, 9),)), (1, 'missing', ((0, 0, 10),))]
                expected = projected_choices(groups, {'A': states}, steps, {'C2'})
                actual = store.states.choices(groups, {'C2'})
                self.assertEqual([list(v) for v in actual], expected)
                self.assertLessEqual(store.states.page.cache_info().currsize, CACHE_PAGES)
                revision = store.first_steps.revision
                store.first_steps[('A', 1)] = 0
                self.assertEqual(store.first_steps.revision, revision + 1)
                self.assertEqual(list(store.states.choices(groups, None)[0])[1][2], 0)
                store.first_steps[('A', 1)] = 0
                self.assertEqual(store.first_steps.revision, revision + 1)
            finally:
                store.close()

    def test_compact_rounds_match_full_rows_states_and_convergence(self):
        structures, reactions, metabolites = fixture()
        sources = ['A[c]', 'B[c]']
        with redirect_stdout(StringIO()):
            expected = list(propagate_rounds('Reference', reactions, metabolites, 8, sources))
        for detail in (0, 1, 3, None):
            with self.subTest(detail=detail), tempfile.TemporaryDirectory() as folder:
                store = self.make_store(folder, detail)
                old_counts, new_counts = {}, {}
                try:
                    with redirect_stdout(StringIO()):
                        actual = propagate_rounds('Disk', reactions, metabolites, 8, sources, row_store=store)
                        for reference in expected:
                            number, _, added, converged = next(actual)
                            self.assertEqual((number, added, converged), (reference[0], reference[2], reference[3]))
                            ids = {name: {name} for name in structures}
                            config = [('A[c]', (), {'C1'}), ('B[c]', (), {'C2'})]
                            self.assertEqual(summarize_sources(store.source_records(), config, ids, number, new_counts),
                                             summarize_sources(reference[1], config, ids, number, old_counts))
                            if detail is not None and number > detail:
                                self.assertEqual(store.db.execute('SELECT COUNT(*) FROM records WHERE round_no=?', (number,)).fetchone()[0], 0)
                                self.assertEqual(len(store.buffer), 0)
                                # Minimum steps and canonical labels still match every emitted product.
                                for row in reference[1].itertuples(index=False):
                                    self.assertEqual(store.catalog[(row.Production, row.LabeledSMILES)], row.ProductionID)
                                    self.assertLessEqual(store.first_steps[(row.Production, row.ProductionID)], row.Step)
                        with self.assertRaises(StopIteration):
                            next(actual)
                    self.assertGreater(len(expected), 2)
                finally:
                    store.close()

    def test_compact_path_does_not_build_detail_rows(self):
        structures, reactions, metabolites = fixture()
        original = FlowOutput.append
        def detail_only(store, row):
            self.assertFalse(store.compact)
            return original(store, row)
        with tempfile.TemporaryDirectory() as folder, patch.object(FlowOutput, 'append', detail_only), redirect_stdout(StringIO()):
            path = Path(folder)/'flow.xlsx'
            trace_analysis('Test', dict(carbon_sources=['A[c]', 'B[c]'], export_detail_rounds=1),
                           reactions, metabolites, 8, Workbook(), path)
            book = load_workbook(path, read_only=True)
            try:
                rows = list(book['Carbon_Flow'].values)
                self.assertTrue(all(row[0] == 1 for row in rows[1:]))
                self.assertGreater(len(list(book['Summary'].values)), 3)
            finally:
                book.close()

    def test_compact_batch_keeps_minimum_step_without_provenance(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.make_store(folder, detail=0, batch=2)
            try:
                store.start_round(4)
                store.append_state('A', 1, '[CH4:1]', 'C1', 9)
                store.append_state('A', 1, '[CH4:1]', 'C1', 4)
                store.flush()
                store.append_state('A', 1, '[CH4:1]', 'C1', 7)
                store.flush()
                self.assertEqual(store.db.execute('SELECT step FROM round_states').fetchone()[0], 4)
                self.assertEqual(store.db.execute('SELECT COUNT(*) FROM records').fetchone()[0], 0)
            finally:
                store.close()

    def test_state_merge_preserves_snapshot_until_boundary(self):
        structures, reactions, _ = fixture()
        with tempfile.TemporaryDirectory() as folder:
            store = self.make_store(folder, detail=0)
            try:
                store.start_round(4)
                store.append_state('P[c]', 1, '[CH3:2][CH2:1]O', 'C1,C2', 2)
                store.flush()
                self.assertIsNone(store.states.get('P[c]'))
                store.merge_states(store.states, store.first_steps, set(), structures)
                self.assertEqual(len(store.states['P[c]']), 1)
                self.assertEqual(store.first_steps[('P[c]', 1)], 2)
            finally:
                store.close()

    def test_workbook_reservation_uses_uncompressed_bytes_before_reading(self):
        from zipfile import ZipFile, ZIP_DEFLATED
        from carbon_flow_tools.config import MemoryGuard
        from carbon_flow_tools.pipeline import read_input
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'compressed.xlsx'
            with ZipFile(str(path), 'w', ZIP_DEFLATED) as archive:
                archive.writestr('xl/worksheets/sheet1.xml', 'x' * 100000)
            with patch.object(MemoryGuard, 'check') as check:
                MemoryGuard().check_file(path)
                self.assertEqual(check.call_args[1]['extra_bytes'], 1200000)
            with patch('carbon_flow_tools.config.available_memory', return_value=(.5, 8)), \
                 patch('carbon_flow_tools.pipeline.pd.ExcelFile') as reader:
                with self.assertRaises(MemoryPressureError):
                    read_input(path, 'Nonzero', 'Metabolites')
                reader.assert_not_called()

    def test_worker_memory_stop_is_not_a_partial_success(self):
        from carbon_flow_tools.enumeration import _candidate_batch
        from carbon_flow_tools.mapper_worker import checked_batch
        from types import SimpleNamespace
        adapter = SimpleNamespace(map_batch=lambda *args: self.fail('prediction should not start'))
        with patch('carbon_flow_tools.config.available_memory', return_value=(.5, 8)):
            with self.assertRaises(MemoryPressureError):
                _candidate_batch([], [], [], {}, 0, False, False, False)
            with self.assertRaises(MemoryPressureError):
                checked_batch(adapter, None, ['C>>C'])

    def test_mapper_does_not_retry_memory_pressure_as_an_ordinary_batch_failure(self):
        import json
        from types import SimpleNamespace
        from unittest.mock import Mock
        from carbon_flow_tools import mapper_worker
        adapter = SimpleNamespace(load_model=lambda: None,
                                  map_batch=Mock(side_effect=MemoryPressureError('fixture pressure')))
        with tempfile.TemporaryDirectory() as folder:
            request, response = Path(folder)/'in.json', Path(folder)/'out.json'
            request.write_text(json.dumps(dict(reactions=['C>>C'])))
            args = ['worker', 'fixture.adapter', str(request), str(response), '1', 'Test', 'Fixture']
            with patch('sys.argv', args), patch.object(mapper_worker.importlib, 'import_module', return_value=adapter), redirect_stdout(StringIO()):
                with self.assertRaises(MemoryPressureError):
                    mapper_worker.main()
            self.assertEqual(adapter.map_batch.call_count, 1)
            self.assertFalse(response.exists())

    def test_buffers_are_bounded_by_bytes_even_with_a_large_row_setting(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.make_store(folder, detail=0, batch=1000000)
            try:
                store.start_round(4)
                for number in range(10):
                    store.append_state('A', number, 'x' * 300000, 'C1', 2)
                    self.assertLess(store.buffer_bytes, 8 * 1024 ** 2)
                self.assertGreater(store.db.execute('SELECT COUNT(*) FROM round_states').fetchone()[0], 0)
            finally:
                store.close()

    def test_only_earlier_steps_keep_a_round_from_converging(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.make_store(folder, detail=0)
            try:
                store.start_round(4)
                store.first_steps[('A', 1)] = 7
                revision = store.first_steps.revision
                store.append_state('A', 1, '[CH4:1]', 'C1', 3)
                store.merge_states(store.states, store.first_steps, {'A'}, {'A': 'C'})
                self.assertEqual(store.first_steps[('A', 1)], 3)
                self.assertGreater(store.first_steps.revision, revision)
                self.assertIsNone(store.states.get('A'))
            finally:
                store.close()

    def test_low_disk_preserves_last_complete_output(self):
        _, reactions, metabolites = fixture()
        original, saved = FlowOutput.complete_round, []
        def complete(store):
            original(store)
            saved.append(store.output.read_bytes())
        original_start = FlowOutput.start_round
        def start(store, number):
            original_start(store, number)
            if number == 2:
                raise DiskSpaceLimitError('fixture disk reserve')
        with tempfile.TemporaryDirectory() as folder, patch.object(FlowOutput, 'complete_round', complete), patch.object(FlowOutput, 'start_round', start), redirect_stdout(StringIO()):
            path, status = Path(folder)/'out.xlsx', {}
            with self.assertRaises(DiskSpaceLimitError):
                trace_analysis('Test', dict(carbon_sources=['A[c]', 'B[c]'], _run_status=status),
                               reactions, metabolites, 8, Workbook(), path)
            self.assertEqual(status['completed_rounds'], 1)
            self.assertEqual(path.read_bytes(), saved[0])
            self.assertFalse(list(Path(folder).glob('.carbon-flow-*')))
        with patch('carbon_flow_tools.config.shutil.disk_usage', return_value=(100*1024**3, 100*1024**3, 1)):
            with self.assertRaises(DiskSpaceLimitError):
                DiskGuard('.').check()


if __name__ == '__main__': unittest.main()
