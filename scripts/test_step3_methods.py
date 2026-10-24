"""Verify EMU-only outputs for ordinary/AI versions and all source counts."""
import hashlib
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch
import pandas as pd
from carbon_flow_tools.pipeline import trace_mapping_file, run_carbon_flow
from scripts.test_parallel_carbon import prepare_fixture


class EMUExportTests(unittest.TestCase):
    def test_export_omits_internal_columns_and_retains_values(self):
        internal = ["FluxResidual", "Status", "StructureType", "KnownCarbonCount", "ContributionScope"]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for sources in (["A[e]"], ["A[e]", "X[e]"]):
                for suffix in ("", "_AIGenerated"):
                    table = pd.DataFrame([dict(MetaboliteID="P", MetaboliteName="P[c]", CarbonCount=float("nan"),
                        Inflow=1, Outflow=1, FluxResidual=0, Status="partial_unknown", StructureType="partial",
                        KnownCarbonCount=2, ContributionScope="explicit carbon positions only", Unknown=.25,
                        **{n: .75 / len(sources) for n in sources})])
                    before = table.copy(deep=True)
                    output = root / ("case" + suffix + ".xlsx")
                    with patch("carbon_flow_tools.pipeline.load_final_mappings", return_value=([], pd.DataFrame(), pd.DataFrame())), patch(
                            "carbon_flow_tools.source_contributions.solve_source_contributions",
                            return_value=({"MetaboliteContributions": table, "BoundaryInputs": pd.DataFrame(columns=["Source"])},
                                          dict(carbon_positions=2))):
                        trace_mapping_file("Case", dict(carbon_sources=sources, min_available_memory_gb=.01),
                                           root / ("Case_2_atom_mapping" + suffix + ".xlsx"), output)
                    actual = pd.read_excel(output, sheet_name="EMU_Metabolites", engine="openpyxl")
                    pd.testing.assert_frame_equal(actual, table.drop(columns=internal), check_dtype=False)
                    pd.testing.assert_frame_equal(table, before)

    def test_versions_keep_distinct_inputs_and_only_emu_sheets(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            normal = prepare_fixture(root)
            ai = root / "Test_2_atom_mapping_AIGenerated.xlsx"
            shutil.copy2(str(normal), str(ai))
            # The AI mapping deliberately leaves one product carbon unresolved.
            from openpyxl import load_workbook
            from carbon_flow_tools.pipeline import load_final_mappings
            from carbon_flow_tools.consensus import carbon_candidate
            reaction = next(r for r in load_final_mappings(normal, {})[0] if r.reaction_id == "r2")
            reaction.mapped_smiles = "[CH3:1]CO>>[CH3:1]CO"
            partial = carbon_candidate(reaction).mapped_smiles
            wb = load_workbook(ai)
            ws = wb["Final_Mapping"]
            fields = [c.value for c in ws[1]]
            for row in range(2, ws.max_row + 1):
                if ws.cell(row, fields.index("ReactionID") + 1).value == "r2":
                    ws.cell(row, fields.index("MappedReactionSMILES") + 1,
                            partial)
                    ws.cell(row, fields.index("MappingCoverage") + 1, "Partial")
            wb.save(ai)
            wb.close()
            hashes = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in (normal, ai)}
            jobs = {"Test": dict(mapping_files=[normal, ai], carbon_sources=["A[e]"], min_available_memory_gb=.01)}
            result = root / "output"
            statuses = run_carbon_flow(jobs, result, max_processes=1)
            self.assertEqual({s["emu"]["mapping_version"] for s in statuses}, {"Original", "AI"})
            for suffix, expected in (("", 1), ("_AIGenerated", .5)):
                with pd.ExcelFile(result / ("Test_3_carbon_flow" + suffix + ".xlsx"), engine="openpyxl") as excel:
                    self.assertEqual(set(excel.sheet_names), {"EMU_Metabolites", "EMU_BoundaryInputs"})
                    table = pd.read_excel(excel, sheet_name="EMU_Metabolites").set_index("MetaboliteName")
                    self.assertAlmostEqual(table.loc["B[c]", "A[e]"], expected)
                    self.assertAlmostEqual(table.loc["B[c]", "Unknown"], 1 - expected)
            self.assertFalse((result / "carbon_flow_summary.xlsx").exists())
            self.assertEqual(hashes, {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in hashes})


if __name__ == "__main__":
    unittest.main()
