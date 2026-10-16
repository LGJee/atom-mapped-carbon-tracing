"""Hand-solvable checks for the optional single-carbon origin balances."""

import unittest

import numpy as np
import pandas as pd

from carbon_flow_tools.reactions import prepare_reactions
from carbon_flow_tools.source_contributions import solve_source_contributions


def solve(structures, specifications, sources):
    metabolites = pd.DataFrame([dict(MetaboliteID=n, MetaboliteName=n, SMILES_Manual=s)
                                for n, s in structures.items()])
    table = pd.DataFrame([dict(ReactionID="r" + str(i), ReactionName="", Flux=v, Equation=e)
                          for i, (e, v, m) in enumerate(specifications)])
    reactions = prepare_reactions(table, metabolites)
    for r, (_, _, mapped) in zip(reactions, specifications):
        r.status = "mapped" if mapped else "missing"
        r.mapped_smiles = mapped
    return solve_source_contributions(reactions, metabolites, sources)


class SourceContributionTests(unittest.TestCase):
    def test_flux_weighted_merge_and_partial_unknown(self):
        for mapping, expected, unknown in [("[CH4:1]>>[CH4:1]", .6, 0), ("", 0, .6)]:
            tables, info = solve({"G": "C", "X": "C", "P": "C"}, [
                (" -> G", 6, ""), (" -> X", 4, ""),
                ("G -> P", 6, mapping), ("X -> P", 4, "[CH4:1]>>[CH4:1]"),
                ("P -> ", 10, "")], ["G", "X"])
            p = tables["MetaboliteContributions"].set_index("MetaboliteID").loc["P"]
            self.assertAlmostEqual(p.G, expected)
            self.assertAlmostEqual(p.X, .4)
            self.assertAlmostEqual(p.Unknown, unknown)
            self.assertLess(info["max_scaled_residual"], 1e-10)

    def test_position_specific_condensation_and_cleavage(self):
        tables, _ = solve({"G": "C", "X": "CO", "P": "CCO"}, [
            (" -> G", 2, ""), (" -> X", 2, ""),
            ("G + X -> P", 3, "[CH4:1].[CH3:2][OH:3]>>[CH3:1][CH2:2][OH:3]"),
            ("P -> G + X", 1, "[CH3:1][CH2:2][OH:3]>>[CH4:1].[CH3:2][OH:3]"),
            ("P -> ", 2, "")], ["G", "X"])
        atoms = tables["AtomContributions"].query("MetaboliteID == 'P'")
        np.testing.assert_allclose(atoms.G, [1, 0])
        np.testing.assert_allclose(atoms.X, [0, 1])
        p = tables["MetaboliteContributions"].set_index("MetaboliteID").loc["P"]
        self.assertAlmostEqual(p.G, .5)
        self.assertAlmostEqual(p.X, .5)

    def test_circulation_does_not_create_extra_source(self):
        tables, _ = solve(dict.fromkeys(["G", "X", "A", "B"], "C"), [
            (" -> G", 6, ""), (" -> X", 4, ""),
            ("G -> A", 6, "[CH4:1]>>[CH4:1]"),
            ("X -> A", 4, "[CH4:1]>>[CH4:1]"),
            ("A -> B", 15, "[CH4:1]>>[CH4:1]"),
            ("B -> A", 5, "[CH4:1]>>[CH4:1]"), ("B -> ", 10, "")], ["G", "X"])
        p = tables["MetaboliteContributions"].set_index("MetaboliteID")
        np.testing.assert_allclose(p.loc[["A", "B"], "G"], [.6, .6])

    def test_unanchored_cycle_and_downstream(self):
        tables, info = solve(dict.fromkeys(["G", "A", "B"], "C"), [
            (" -> G", 1, ""), ("G -> ", 1, ""),
            ("A -> B", 3, "[CH4:1]>>[CH4:1]"),
            ("B -> A", 3, "[CH4:1]>>[CH4:1]")], ["G"])
        p = tables["MetaboliteContributions"].set_index("MetaboliteID")
        self.assertEqual(info["unanchored_positions"], 2)
        self.assertEqual(p.loc["A", "Status"], "unanchored_cycle")
        self.assertAlmostEqual(p.loc["A", "Unknown"], 1)

    def test_negative_flux_and_repeated_molecules(self):
        tables, _ = solve({"G": "C", "P": "C"}, [
            ("G -> ", -2, ""),
            ("2 P -> 2 G", -1, "[CH4:1].[CH4:2]>>[CH4:1].[CH4:2]"),
            ("P -> ", 2, "")], ["G"])
        p = tables["MetaboliteContributions"].set_index("MetaboliteID")
        self.assertAlmostEqual(p.loc["P", "G"], 1)
        self.assertEqual(p.loc["P", "Status"], "complete")

    def test_external_other_and_unbalanced_pool(self):
        for output_rate in (1, 2):
            tables, _ = solve(dict.fromkeys(["G", "Z", "P"], "C"), [
                (" -> G", 1, ""), ("G -> ", 1, ""),
                (" -> Z", 1, ""), ("Z -> P", 1, "[CH4:1]>>[CH4:1]"),
                ("P -> ", output_rate, "")], ["G"])
            p = tables["MetaboliteContributions"].set_index("MetaboliteID")
            if output_rate == 1:
                self.assertAlmostEqual(p.loc["P", "External:Z"], 1)
            else:
                self.assertEqual(p.loc["P", "Status"], "invalid_flux_balance")
                self.assertAlmostEqual(p.loc["P", "Unknown"], 1)

    def test_same_origin_symmetry_preserves_known_source(self):
        tables, info = solve({"G": "CC", "P": "CC"}, [
            (" -> G", 1, ""), ("G -> P", 1, "[CH3:1][CH3:2]>>[CH3:1][CH3:2]"),
            ("P -> ", 1, "")], ["G"])
        p = tables["MetaboliteContributions"].set_index("MetaboliteID")
        self.assertEqual(info["ambiguous_carbon_positions"], 4)
        self.assertAlmostEqual(p.loc["P", "G"], 1)
        self.assertAlmostEqual(p.loc["P", "Unknown"], 0)

    def test_same_origin_survives_symmetric_intermediate_and_cleavage(self):
        tables, info = solve({"G": "CCCCO", "P": "CC(C)CO", "Q": "C", "R": "CCCO"}, [
            (" -> G", 1, ""),
            ("G -> P", 1, "[CH3:1][CH2:2][CH2:3][CH2:4][OH:5]>>[CH3:1][CH:2]([CH3:3])[CH2:4][OH:5]"),
            ("P -> Q + R", 1, "[CH3:1][CH:2]([CH3:3])[CH2:4][OH:5]>>[CH4:1].[CH3:3][CH2:2][CH2:4][OH:5]"),
            ("Q -> ", 1, ""), ("R -> ", 1, "")], ["G"])
        p = tables["MetaboliteContributions"].set_index("MetaboliteID")
        np.testing.assert_allclose(p.loc[["P", "Q", "R"], "G"], 1)
        np.testing.assert_allclose(p.loc[["P", "Q", "R"], "Unknown"], 0)
        self.assertGreater(info["symmetry_mapped_transfers"], 0)

    def test_symmetric_positions_follow_the_selected_mapping_without_scrambling(self):
        for cleavage, expected_g in (("[CH3:1][CH3:2]>>[CH4:1].[CH4:2]", [1, 0]),
                                     ("[CH3:1][CH3:2]>>[CH4:2].[CH4:1]", [0, 1])):
            tables, _ = solve({"G": "C", "X": "C", "P": "CC", "Q": "C", "R": "C"}, [
                (" -> G", 1, ""), (" -> X", 1, ""),
                ("G + X -> P", 1, "[CH4:1].[CH4:2]>>[CH3:1][CH3:2]"),
                ("P -> Q + R", 1, cleavage),
                ("Q -> ", 1, ""), ("R -> ", 1, "")], ["G", "X"])
            p = tables["MetaboliteContributions"].set_index("MetaboliteID")
            self.assertAlmostEqual(p.loc["P", "G"], .5)
            self.assertAlmostEqual(p.loc["P", "X"], .5)
            np.testing.assert_allclose(p.loc[["Q", "R"], "G"], expected_g)
            np.testing.assert_allclose(p.loc[["P", "Q", "R"], "Unknown"], 0)

    def test_symmetric_transfer_preserves_partial_mixed_origins(self):
        tables, _ = solve(dict.fromkeys(["G", "X", "Z", "P", "Q"], "CC"), [
            (" -> G", 4, ""), (" -> X", 2, ""), (" -> Z", 4, ""),
            ("G -> P", 4, "[CH3:1][CH3:2]>>[CH3:1][CH3:2]"),
            ("X -> P", 2, "[CH3:1][CH3:2]>>[CH3:1][CH3:2]"),
            ("Z -> P", 4, ""),
            ("P -> Q", 10, "[CH3:1][CH3:2]>>[CH3:1][CH3:2]"),
            ("Q -> ", 10, "")], ["G", "X"])
        p = tables["MetaboliteContributions"].set_index("MetaboliteID")
        for name in ("P", "Q"):
            np.testing.assert_allclose(p.loc[name, ["G", "X", "Unknown"]].astype(float), [.4, .2, .4])

    def test_two_precursor_copies_preserve_mean_origins_in_a_symmetric_product(self):
        structures = dict.fromkeys(["G", "X", "Z", "P"], "CCO")
        structures["Q"] = "OCCCCO"
        tables, _ = solve(structures, [
            (" -> G", 4, ""), (" -> X", 2, ""), (" -> Z", 4, ""),
            ("G -> P", 4, "[CH3:1][CH2:2][OH:3]>>[CH3:1][CH2:2][OH:3]"),
            ("X -> P", 2, "[CH3:1][CH2:2][OH:3]>>[CH3:1][CH2:2][OH:3]"),
            ("Z -> P", 4, ""),
            ("2 P -> Q", 5, "[CH3:1][CH2:2][OH:3].[CH3:4][CH2:5][OH:6]>>[OH:3][CH2:2][CH2:1][CH2:4][CH2:5][OH:6]"),
            ("Q -> ", 5, "")], ["G", "X"])
        p = tables["MetaboliteContributions"].set_index("MetaboliteID")
        np.testing.assert_allclose(p.loc["Q", ["G", "X", "Unknown"]].astype(float), [.4, .2, .4])

    def test_unknown_precursor_is_not_resolved_just_because_one_source_is_configured(self):
        tables, _ = solve({"G": "CC", "P": "CC", "Q": "CC"}, [
            (" -> G", 1, ""), ("G -> P", 1, ""),
            ("P -> Q", 1, "[CH3:1][CH3:2]>>[CH3:1][CH3:2]"),
            ("Q -> ", 1, "")], ["G"])
        p = tables["MetaboliteContributions"].set_index("MetaboliteID")
        self.assertAlmostEqual(p.loc["Q", "Unknown"], 1)

    def test_unanchored_position_in_symmetric_recycling_stays_unknown(self):
        tables, _ = solve({"G": "C", "P": "CC", "W": "C"}, [
            (" -> G", 1, ""),
            ("G + P -> P + W", 1, "[CH4:1].[CH3:2][CH3:3]>>[CH3:1][CH3:2].[CH4:3]"),
            ("W -> ", 1, "")], ["G"])
        p = tables["MetaboliteContributions"].set_index("MetaboliteID")
        self.assertAlmostEqual(p.loc["P", "G"], .5)
        self.assertAlmostEqual(p.loc["P", "Unknown"], .5)
        self.assertAlmostEqual(p.loc["W", "Unknown"], 1)

    def test_identical_inflow_composition_at_configured_source_pool(self):
        tables, _ = solve(dict.fromkeys(["G", "X"], "C"), [
            (" -> G", 1, ""), (" -> X", 1, ""),
            ("G -> X", 1, "[CH4:1]>>[CH4:1]"),
            ("X -> ", 2, "")], ["G", "X"])
        p = tables["MetaboliteContributions"].set_index("MetaboliteID")
        self.assertAlmostEqual(p.loc["X", "G"], .5)
        self.assertAlmostEqual(p.loc["X", "X"], .5)

    def test_scientific_notation_stoichiometry(self):
        from carbon_flow_tools.reactions import parse_side
        self.assertEqual(parse_side("2.5e-05 Co2+[c] + .5 A[c] + 2E+1 B[c]"),
                         [(2.5e-5, "Co2+[c]"), (.5, "A[c]"), (20., "B[c]")])

    def test_each_position_can_have_multiple_origins(self):
        tables, _ = solve({"G": "C", "X": "CO", "P": "CCO"}, [
            (" -> G", 10, ""), (" -> X", 10, ""),
            ("G + X -> P", 6, "[CH4:1].[CH3:2][OH:3]>>[CH3:1][CH2:2][OH:3]"),
            ("G + X -> P", 4, "[CH4:1].[CH3:2][OH:3]>>[CH3:2][CH2:1][OH:3]"),
            ("P -> ", 10, "")], ["G", "X"])
        atoms = tables["AtomContributions"].query("MetaboliteID == 'P'")
        np.testing.assert_allclose(atoms.G, [.6, .4])
        np.testing.assert_allclose(atoms.X, [.4, .6])

    def test_structure_formula_mismatch_is_unavailable(self):
        structures = pd.DataFrame([dict(MetaboliteID="G", MetaboliteName="G", SMILES_Manual="C"),
            dict(MetaboliteID="P", MetaboliteName="P", SMILES_Manual="CO", MolecularFormula="C2H6O")])
        source = pd.DataFrame([dict(ReactionID="in", ReactionName="", Flux=1, Equation=" -> G"),
            dict(ReactionID="out", ReactionName="", Flux=1, Equation="G -> ")])
        tables, _ = solve_source_contributions(prepare_reactions(source, structures), structures, ["G"])
        p = tables["MetaboliteContributions"].set_index("MetaboliteID").loc["P"]
        self.assertEqual(p.Status, "Structure/formula carbon count mismatch")
        self.assertTrue(pd.isna(p.G))

    def test_workbook_entry_accepts_both_mapping_versions_and_preserves_inputs(self):
        from hashlib import sha256
        import json
        from pathlib import Path
        import tempfile
        from openpyxl import Workbook
        from carbon_flow_tools.excel_io import write_sheet, save_output
        from analyze_source_contributions import main
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "Tiny_1_input.xlsx"
            mapping = root / "Tiny_2_atom_mapping.xlsx"
            ai = root / "Tiny_2_atom_mapping_AIGenerated.xlsx"
            rows = pd.DataFrame([dict(ReactionID="in", ReactionName="", Flux=1, Equation=" -> G"),
                dict(ReactionID="r", ReactionName="", Flux=1, Equation="G -> P"),
                dict(ReactionID="out", ReactionName="", Flux=1, Equation="P -> ")])
            metabolites = pd.DataFrame([dict(MetaboliteID=n, MetaboliteName=n, SMILES_Manual="C") for n in ("G", "P")])
            w = Workbook()
            del w[w.active.title]
            write_sheet(w, "Nonzero", rows)
            write_sheet(w, "Metabolites", metabolites)
            save_output(w, source)
            digest = sha256(source.read_bytes()).hexdigest()
            write_sheet(w, "Mapping_Info", pd.DataFrame([dict(SourceInputSHA256=digest)]))
            write_sheet(w, "Final_Mapping", pd.DataFrame([dict(ReactionID="r", Equation="G -> P",
                MappingDirection="forward", MappedReactionSMILES="[CH4:1]>>[CH4:1]")]))
            save_output(w, mapping)
            w["Final_Mapping"].cell(1, 5, "AIGenerated")
            w["Final_Mapping"].cell(2, 5, True)
            w["Final_Mapping"].cell(1, 6, "AIReviewed")
            w["Final_Mapping"].cell(2, 6, True)
            save_output(w, ai)
            w.close()
            before = {p: sha256(p.read_bytes()).hexdigest() for p in (source, mapping, ai)}
            plan = root / "plan.json"
            plan.write_text(json.dumps(dict(version=3, output_dir=str(root), models=dict(Tiny=dict(
                input_file=str(source), input_sha256=digest, carbon_sources=["G"])))), encoding="utf-8")
            self.assertEqual(main(["--plan", str(plan), "--output-dir", str(root / "output")]), 0)
            for p, value in before.items():
                self.assertEqual(sha256(p.read_bytes()).hexdigest(), value)
            run = json.loads((root / "output/source_contributions_run.json").read_text())
            self.assertEqual(run["status"], "finished")
            self.assertEqual(len(run["models"]), 2)
            self.assertEqual({m["mapping_version"] for m in run["models"]}, {"Original", "AI"})
            for m in run["models"]:
                d = pd.read_excel(m["output_file"], sheet_name="MetaboliteContributions", engine="openpyxl")
                self.assertTrue(d.G.eq(1).all())
                self.assertTrue(d.Unknown.eq(0).all())
            changed = rows.copy()
            changed.loc[0, "Flux"] = 2
            w = Workbook()
            del w[w.active.title]
            write_sheet(w, "Nonzero", changed)
            write_sheet(w, "Metabolites", metabolites)
            write_sheet(w, "Mapping_Info", pd.DataFrame([dict(SourceInputSHA256=digest)]))
            write_sheet(w, "Final_Mapping", pd.DataFrame([dict(ReactionID="r", Equation="G -> P",
                MappingDirection="forward", MappedReactionSMILES="[CH4:1]>>[CH4:1]")]))
            save_output(w, mapping)
            w.close()
            with self.assertRaises(AssertionError):
                main(["--plan", str(plan), "--output-dir", str(root / "rejected")])
            self.assertEqual(json.loads((root / "rejected/source_contributions_run.json").read_text())["status"], "error")

    def test_unavailable_positions_do_not_break_mapped_transfers(self):
        rows = pd.DataFrame([
            dict(ReactionID="in", ReactionName="", Flux=1, Equation=" -> G"),
            dict(ReactionID="r", ReactionName="", Flux=1, Equation="G -> P"),
            dict(ReactionID="out", ReactionName="", Flux=1, Equation="P -> ")])
        for unavailable in ({"G"}, {"P"}, {"G", "P"}):
            with self.subTest(unavailable=unavailable):
                structures = pd.DataFrame([dict(MetaboliteID=n, MetaboliteName=n,
                    SMILES_Manual="C", MolecularFormula="C2" if n in unavailable else "CH4")
                    for n in ("G", "P")])
                reactions = prepare_reactions(rows, structures)
                reactions[1].mapped_smiles = "[CH4:1]>>[CH4:1]"
                reactions[1].status = "mapped"
                # A separate usable source keeps this case's other inputs unchanged.
                structures = pd.concat([structures, pd.DataFrame([dict(MetaboliteID="X", MetaboliteName="X",
                    SMILES_Manual="C", MolecularFormula="CH4")])], ignore_index=True)
                tables, _ = solve_source_contributions(reactions, structures, ["X"])
                result = tables["MetaboliteContributions"].set_index("MetaboliteID")
                if "P" in unavailable:
                    self.assertTrue(pd.isna(result.loc["P", "Unknown"]))
                    self.assertEqual(result.loc["P", "Status"], "Structure/formula carbon count mismatch")
                else:
                    self.assertAlmostEqual(result.loc["P", "Unknown"], 1)
                for name in unavailable:
                    self.assertFalse(tables["AtomContributions"].MetaboliteID.eq(name).any())

    def test_invalid_mapping_is_rejected(self):
        with self.assertRaises(ValueError):
            solve({"G": "C", "P": "C"}, [(" -> G", 1, ""),
                ("G -> P", 1, "[CH4:1]>>[CH4:2]"), ("P -> ", 1, "")], ["G"])


if __name__ == "__main__":
    unittest.main()
