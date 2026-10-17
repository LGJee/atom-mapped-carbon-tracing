"""Check source exclusivity across all recorded states and final-round exports."""
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook

from carbon_flow_tools.excel_io import WorkbookSizeLimitError
from carbon_flow_tools.flow_output import FlowOutput
from carbon_flow_tools.pipeline import summarize_sources, trace_analysis
from scripts.test_output_limits import flow_fixture


class ExclusiveSourceTests(unittest.TestCase):
    def test_all_states_must_belong_to_one_source(self):
        sources = [("Glucose", (), {"C1", "C2"}),
                   ("Xylose", (), {"C3", "C4"}), ("Other", (), {"C5"})]
        records = [("pure[c]", "C1"), ("pure[c]", "C2,C2"),
                   ("pure[c]", "C1"), ("pure[e]", "C3"),
                   ("separate", "C1"), ("separate", "C3"),
                   ("mixed", "C1,C3"), ("third", "C5"),
                   ("unknown", "C99"), ("empty", ""),
                   ("invalid", "C1"), ("invalid", "C99"),
                   ("not_mapped", "C1")]
        ids = {name: {name} for name, _ in records if name != "not_mapped"}
        ids["pure[c]"] = {"id1", "id2"}
        details = []
        row = summarize_sources(iter(records), sources, ids, 3, {}, details)
        self.assertEqual([row[name + "_ExclusiveSourceMetabolites"]
                          for name, _, _ in sources], [2, 1, 1])
        self.assertEqual(details, [(3, "Glucose", "id1", "pure[c]"),
                                  (3, "Glucose", "id2", "pure[c]"),
                                  (3, "Xylose", "pure[e]", "pure[e]"),
                                  (3, "Other", "third", "third")])
        self.assertEqual(row["MixedSources_LabeledCarbonMetabolites"], 1)
        frame = pd.DataFrame(records, columns=["Production", "CarbonLabels"])
        self.assertEqual(row, summarize_sources(frame, sources, ids, 3, {}))
        single = summarize_sources(iter(records), sources[:1], ids, 3, {})
        self.assertEqual(single["Glucose_ExclusiveSourceMetabolites"], 2)

    def test_only_last_saved_round_has_counts_and_optional_list(self):
        rxn, metabolites, _ = flow_fixture()
        for suffix in ("", "_AIGenerated"):
            for export in (False, True):
                with self.subTest(suffix=suffix, export=export), tempfile.TemporaryDirectory() as directory:
                    output = Path(directory) / ("Test_3_carbon_flow" + suffix + ".xlsx")
                    trace_analysis("Test", {"carbon_sources": ["A[c]"],
                        "export_exclusive_source_metabolites": export},
                        [rxn], metabolites, 3, Workbook(), output)
                    workbook = load_workbook(output, read_only=True)
                    try:
                        rows = list(workbook["Summary"].values)
                        col = rows[0].index("A[c]_ExclusiveSourceMetabolites")
                        self.assertGreater(len(rows), 2)
                        self.assertTrue(all(row[col] is None for row in rows[1:-1]))
                        self.assertEqual(rows[-1][col], 1)
                        sheet = "Exclusive_Source_Metabolites"
                        self.assertEqual(sheet in workbook.sheetnames, export)
                        if export:
                            self.assertEqual(list(workbook[sheet].values), [
                                ("Round", "CarbonSource", "MetaboliteID", "MetaboliteName"),
                                (rows[-1][0], "A[c]", "B[c]", "B[c]")])
                    finally:
                        workbook.close()

    def test_later_alternative_source_removes_exclusivity(self):
        rxn, metabolites, table = flow_fixture()
        metabolites = pd.concat([metabolites, pd.DataFrame([
            dict(MetaboliteID="X[c]", MetaboliteName="X[c]", SMILES_Model="C")])],
            ignore_index=True)
        def rounds(*args, **kwargs):
            store = kwargs["row_store"]
            row = table.to_dict("records")[0]
            for number in (1, 2):
                store.start_round(number)
                store.start_layer()
                store.append(row)
                if number == 2:
                    store.append(dict(row, CarbonLabels="C2", ProductionID=2))
                store.flush()
                yield number, None, 1, number == 2
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "Test_3_carbon_flow.xlsx"
            with patch("carbon_flow_tools.pipeline.propagate_rounds", rounds):
                trace_analysis("Test", {"carbon_sources": ["A[c]", "X[c]"],
                    "export_exclusive_source_metabolites": True},
                    [rxn], metabolites, 2, Workbook(), output)
            workbook = load_workbook(output, read_only=True)
            try:
                rows = list(workbook["Summary"].values)
                for source in ("A[c]", "X[c]"):
                    col = rows[0].index(source + "_ExclusiveSourceMetabolites")
                    self.assertIsNone(rows[1][col])
                    self.assertEqual(rows[2][col], 0)
                self.assertEqual(len(list(workbook["Exclusive_Source_Metabolites"].values)), 1)
            finally:
                workbook.close()

    def test_rejected_round_preserves_counts_and_list_for_previous_round(self):
        rxn, metabolites, _ = flow_fixture()
        complete = FlowOutput.complete_round
        def reject_second(store):
            if store.number == 2:
                raise WorkbookSizeLimitError("simulated size limit")
            complete(store)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "Test_3_carbon_flow.xlsx"
            status = {}
            with patch.object(FlowOutput, "complete_round", reject_second):
                trace_analysis("Test", {"carbon_sources": ["A[c]"],
                    "export_exclusive_source_metabolites": True, "_run_status": status},
                    [rxn], metabolites, 3, Workbook(), output)
            self.assertEqual(status["completed_rounds"], 1)
            self.assertEqual(status["stop_reason"], "file_size_limit")
            workbook = load_workbook(output, read_only=True)
            try:
                rows = list(workbook["Summary"].values)
                self.assertEqual(len(rows), 2)
                self.assertEqual(rows[1][rows[0].index("A[c]_ExclusiveSourceMetabolites")], 1)
                self.assertEqual(list(workbook["Exclusive_Source_Metabolites"].values)[1][0], 1)
                self.assertNotIn("Carbon_Flow_R2", workbook.sheetnames)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
