"""Check concise output, live ETA revisions and actionable input errors."""

import subprocess
import sys
import json
from contextlib import redirect_stderr, redirect_stdout
from io import StringIO
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from carbon_flow_tools import cli
from carbon_flow_tools.console import Progress, detail
from carbon_flow_tools.pipeline import read_input
from scripts.test_parallel_carbon import make_input


def load_model(options):
    import warnings
    import logging
    warnings.warn("fixture dependency deprecation", UserWarning)
    logging.warning("fixture token sequence length warning")
    os.write(2, b"fixture native chemistry warning\n")
    print("fixture library startup message")
    if options.get("fail"):
        raise RuntimeError("fixture load failed")
    return None


def map_batch(model, reactions):
    import logging
    logging.warning("fixture prediction warning after model load")
    print("fixture library prediction message")
    return [{"mapped_rxn": "[CH4:1]>>[CH4:1]"} if value == "C>>C"
            else {"error": "fixture no prediction"} for value in reactions]


class ConsoleTests(unittest.TestCase):
    def test_eta_is_revised_when_recent_processing_slows(self):
        now = [0.0]
        output = StringIO()
        with patch.dict(os.environ, {"CARBON_TRACING_VERBOSE": "1"}), patch("carbon_flow_tools.console.time.monotonic", side_effect=lambda: now[0]), redirect_stdout(output):
            progress = Progress("Mapper", 100, "reactions")
            progress.update(force=True)
            now[0] = 15
            progress.update(10)
            first_rate = progress.rate
            now[0] = 30
            progress.update(12)
            self.assertLess(progress.rate, first_rate)
            now[0] = 45
            progress.update(100, force=True)
        lines = output.getvalue().splitlines()
        self.assertIn("Estimating...", lines[0])
        self.assertIn("ETA ~2m 15s", lines[1])
        self.assertIn("ETA ~3m 40s", lines[2])
        self.assertIn("100/100 reactions", lines[-1])
        self.assertIn("ETA 0s", lines[-1])

    def test_progress_is_throttled_across_reactions(self):
        now = [0.0]
        output = StringIO()
        with patch("carbon_flow_tools.console.time.monotonic", side_effect=lambda: now[0]), redirect_stdout(output):
            progress = Progress("R1", 100)
            progress.update(force=True)
            now[0] = 5
            progress.begin("R2", 200)
            progress.update(10)
            now[0] = 15
            progress.update(20)
        self.assertEqual(len(output.getvalue().splitlines()), 2)
        self.assertIn("R2", output.getvalue())

    def test_huge_combination_count_cannot_break_a_calculation(self):
        now = [0.0]
        output = StringIO()
        with patch.dict(os.environ, {"CARBON_TRACING_VERBOSE": "1"}), patch("carbon_flow_tools.console.time.monotonic", side_effect=lambda: now[0]), redirect_stdout(output):
            progress = Progress("Reaction", 10 ** 5000, "combinations")
            now[0] = 15
            progress.update(100)
        self.assertIn(">10^300 combinations", output.getvalue())
        self.assertIn("ETA Unavailable", output.getvalue())

    def test_unknown_work_does_not_invent_an_eta(self):
        output = StringIO()
        with patch.dict(os.environ, {"CARBON_TRACING_VERBOSE": "0"}), redirect_stdout(output):
            Progress("Loading database", unit="records").update(25, force=True)
        self.assertIn("still working", output.getvalue())
        self.assertNotIn("ETA", output.getvalue())

    def test_verbose_controls_diagnostics(self):
        output = StringIO()
        with patch.dict(os.environ, {"CARBON_TRACING_VERBOSE": "0"}), redirect_stdout(output):
            detail("hidden")
        self.assertEqual(output.getvalue(), "")
        with patch.dict(os.environ, {"CARBON_TRACING_VERBOSE": "1"}), redirect_stdout(output):
            detail("visible")
        self.assertEqual(output.getvalue(), "visible\n")

    def test_default_progress_uses_plain_language(self):
        output = StringIO()
        with patch.dict(os.environ, {"CARBON_TRACING_VERBOSE": "0"}), redirect_stdout(output):
            progress = Progress("Model / RXNMapper", 10, "reactions")
            progress.update(5, force=True)
            progress.update(10, force=True)
        self.assertIn("50% (5/10 reactions)", output.getvalue())
        self.assertIn("finished (10 reactions)", output.getvalue())
        for jargon in (" | ", "ETA", "Estimating", "Unavailable"):
            self.assertNotIn(jargon, output.getvalue())

    def test_worker_logs_python_native_warnings_and_retains_failures(self):
        root = Path(__file__).resolve().parents[1]
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            request, response, log = [folder / name for name in ("request.json", "response.json", "mapper.log")]
            for verbose, fail in ((False, False), (True, False), (False, True)):
                request.write_text(json.dumps(dict(reactions=["C>>C", "N>>N"],
                    options={"fail":fail}, log_file=str(log))))
                result = subprocess.run([sys.executable, str(root/"carbon_flow_tools/mapper_worker.py"),
                    "scripts.test_console", str(request), str(response), "1", "Example", "TestMapper"],
                    stdout=subprocess.PIPE, stderr=subprocess.PIPE, universal_newlines=True,
                    env=dict(os.environ, CARBON_TRACING_VERBOSE="1" if verbose else "0"))
                self.assertEqual(result.returncode, 1 if fail else 0)
                terminal = result.stdout + result.stderr
                if verbose:
                    self.assertIn("fixture dependency deprecation", terminal)
                    self.assertIn("fixture native chemistry warning", terminal)
                else:
                    self.assertNotIn("fixture native chemistry warning", terminal)
                    self.assertNotIn("UserWarning", terminal)
                    self.assertNotIn("fixture library", terminal)
                    self.assertIn("fixture native chemistry warning", log.read_text())
                    self.assertIn("fixture token sequence length warning", log.read_text())
                    self.assertNotIn("I/O operation on closed file", log.read_text())
                    if not fail:
                        self.assertIn("fixture prediction warning after model load", log.read_text())
                if fail:
                    self.assertIn("Mapper worker failed", terminal)
                    self.assertNotIn("finished", terminal)
                    self.assertIn("Traceback", log.read_text())
                else:
                    self.assertIn("1 reaction needs checking", terminal)
                    self.assertEqual(json.loads(response.read_text()), [
                        {"mapped_rxn":"[CH4:1]>>[CH4:1]"}, {"error":"fixture no prediction"}])
                    self.assertIn("fixture no prediction", log.read_text())

    def test_expected_cli_error_is_concise_and_nonzero(self):
        output = StringIO()
        with patch.object(cli, "main", side_effect=FileNotFoundError("File: missing.xlsx")), redirect_stderr(output):
            result = cli.run_cli("map", {}, [])
        self.assertEqual(result, 1)
        self.assertIn("Step 2 stopped", output.getvalue())
        self.assertIn("missing.xlsx", output.getvalue())
        self.assertIn("Action:", output.getvalue())
        self.assertNotIn("Traceback", output.getvalue())

    def test_verbose_retains_original_exception(self):
        with patch.object(cli, "main", side_effect=ValueError("invalid")):
            with self.assertRaisesRegex(ValueError, "invalid"):
                cli.run_cli("trace", {}, ["--verbose"])

    def test_real_entries_hide_expected_tracebacks_unless_verbose(self):
        root = Path(__file__).resolve().parents[1]
        with tempfile.TemporaryDirectory() as temporary:
            missing = Path(temporary) / "missing-plan.json"
            for entry in ("step2_atom_mapping.py", "step3_carbon_flow.py"):
                for verbose in (False, True):
                    command = [sys.executable, "-B", str(root / entry), "--plan", str(missing)]
                    if verbose:
                        command.append("--verbose")
                    result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                            universal_newlines=True)
                    self.assertNotEqual(result.returncode, 0)
                    self.assertIn(str(missing), result.stderr)
                    self.assertEqual("Traceback" in result.stderr, verbose)

    def test_permission_error_gives_excel_guidance(self):
        output = StringIO()
        with patch.object(cli, "main", side_effect=PermissionError("output.xlsx")), redirect_stderr(output):
            self.assertEqual(cli.run_cli("trace", {}, []), 1)
        self.assertIn("Close the workbook in Excel", output.getvalue())

    def test_workbook_errors_identify_file_sheet_and_rows(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "input.xlsx"
            make_input(path, "B[c]")
            original = path.read_bytes()
            workbook = load_workbook(path)
            workbook["Nonzero"]["A3"] = "r1"
            workbook.save(path)
            workbook.close()
            before = path.read_bytes()
            with self.assertRaises(ValueError) as caught:
                read_input(path, "Nonzero", "Metabolites")
            message = str(caught.exception)
            for text in (str(path), "Sheet: Nonzero", "2 (r1)", "3 (r1)", "Action:"):
                self.assertIn(text, message)
            self.assertEqual(path.read_bytes(), before)
            path.write_bytes(original)
            workbook = load_workbook(path)
            workbook["Nonzero"]["D2"] = 0
            workbook.save(path)
            workbook.close()
            with self.assertRaises(ValueError) as caught:
                read_input(path, "Nonzero", "Metabolites")
            self.assertIn("2 (r1)", str(caught.exception))
            self.assertIn(str(path), str(caught.exception))

    def test_missing_columns_name_the_affected_sheet(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "input.xlsx"
            make_input(path, "B[c]")
            workbook = load_workbook(path)
            workbook["Nonzero"]["B1"] = "wrong"
            workbook.save(path)
            workbook.close()
            with self.assertRaises(ValueError) as caught:
                read_input(path, "Nonzero", "Metabolites")
            self.assertIn("ReactionName", str(caught.exception))
            self.assertIn("Sheet: Nonzero", str(caught.exception))
            self.assertIn(str(path), str(caught.exception))


if __name__ == "__main__":
    unittest.main()
