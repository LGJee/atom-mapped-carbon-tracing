"""Known-structure mapper views retain the original equation and flux."""
from copy import deepcopy
import json
import csv
from pathlib import Path
import tempfile
import unittest
import pandas as pd
from openpyxl import load_workbook
from carbon_flow_tools.reactions import prepare_reactions
from carbon_flow_tools.consensus import select_mappings
from carbon_flow_tools.carbon_mapping_rules import apply_rules, apply_manual_mappings, COLUMNS, validate_manual_records, partial_mapped_smiles
from carbon_flow_tools.carbon_trace import prepare_mappings, map_one_reaction
from carbon_flow_tools.pipeline import finish_analysis, load_final_mappings
from carbon_flow_tools.source_contributions import solve_source_contributions

MAPPED = "[CH3:1][CH2:2][OH:3].[CH3:4][CH2:5][OH:6]>>[CH3:1][CH2:2][OH:3].[CH3:4][CH2:5][OH:6]"

def fixture(flux=1):
    source = pd.DataFrame([["R", "Carrier-dependent", flux,
        "2 A[c] + ReducedProtein[c] => 2 B[c] + OxidizedProtein[c]"]],
        columns=["ReactionID", "ReactionName", "Flux", "Equation"])
    metabolites = pd.DataFrame([["a", "A[c]", "CCO"], ["b", "B[c]", "CCO"],
        ["red", "ReducedProtein[c]", ""], ["ox", "OxidizedProtein[c]", ""]],
        columns=["MetaboliteID", "MetaboliteName", "SMILES_Model"])
    predictions = {name: prepare_reactions(source, metabolites) for name in ("LocalMapper", "RXNMapper", "EnzymeMap")}
    for rows in predictions.values():
        rows[0].mapped_smiles, rows[0].status = MAPPED, "mapped"
    selected, choices = select_mappings(predictions)
    apply_rules(selected, choices, metabolites)
    return source, metabolites, predictions, selected, choices

class ReducedMapperTests(unittest.TestCase):
    def test_forward_reverse_preparation_and_enumeration(self):
        for flux, donor, product in ((1, "A[c]", "B[c]"), (-1, "B[c]", "A[c]")):
            source, metabolites, predictions, selected, choices = fixture(flux)
            r = predictions["EnzymeMap"][0]
            self.assertEqual((r.flux, r.equation), (flux, source.iloc[0].Equation))
            self.assertEqual(len(r.omitted_components), 2)
            self.assertEqual(r.mapping_equation, "2 " + donor + " => 2 " + product)
            self.assertEqual((len(r.left), len(r.right)), (2, 2))
            self.assertEqual(choices.iloc[0].MappingCoverage, "Partial")
            self.assertTrue(choices.iloc[0].ReviewRequired)
            self.assertIn("missing-SMILES participants", choices.iloc[0].SelectionReason)
            self.assertEqual(r.mapped_smiles, MAPPED)
            prepared = prepare_mappings(selected)
            _, rows, _ = map_one_reaction(*prepared[0], {donor: {1: ((0, 0, "C1"), (0, 1, "C2"))}},
                1, {}, 2, dict(zip(metabolites.MetaboliteName, metabolites.SMILES_Model)))
            self.assertTrue(rows)
            self.assertEqual({row["Production"] for row in rows}, {product})
            self.assertTrue(all(row["CarbonLabels"] == "C1,C2" for row in rows))

    def test_missing_side_invalid_structures_and_unknown_names(self):
        source, metabolites, _, _, _ = fixture()
        for equation in ("A[c] => ReducedProtein[c]", "ReducedProtein[c] => A[c]", "A[c] => Absent[c]"):
            item = source.copy()
            item.loc[0, "Equation"] = equation
            r = prepare_reactions(item, metabolites)[0]
            self.assertNotEqual(r.status, "ready")
            self.assertEqual(r.reaction_smiles, "")
        metabolites.loc[metabolites.MetaboliteName.eq("A[c]"), "SMILES_Model"] = "not-a-smiles"
        self.assertIn("Invalid SMILES", prepare_reactions(source, metabolites)[0].status)

    def test_carbon_free_retained_view_is_not_a_carbon_map(self):
        source, metabolites, _, _, _ = fixture()
        metabolites.loc[metabolites.MetaboliteID.isin(["a", "b"]), "SMILES_Model"] = "O"
        r = prepare_reactions(source, metabolites)[0]
        r.mapped_smiles, r.status = "[OH2:1].[OH2:2]>>[OH2:1].[OH2:2]", "mapped"
        selected, choices = select_mappings({"RXNMapper": [r]})
        apply_rules(selected, choices, metabolites)
        self.assertEqual(choices.iloc[0].MappingCoverage, "Unresolved")
        self.assertTrue(choices.iloc[0].SelectedMappedReactionSMILES)

    def test_workbook_roundtrip_and_false_complete_rejection(self):
        source, metabolites, predictions, _, _ = fixture()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "input.xlsx"
            with pd.ExcelWriter(path, engine="openpyxl") as writer:
                source.to_excel(writer, sheet_name="Nonzero", index=False)
                metabolites.to_excel(writer, sheet_name="Metabolites", index=False)
            finish_analysis("Test", dict(input_file=path, metabolite_sheet="Metabolites"), metabolites,
                predictions, [], 1, list(predictions), root, stage="map")
            output = root / "Test_2_atom_mapping.xlsx"
            loaded, saved, _ = load_final_mappings(output, {})
            self.assertEqual(loaded[0].status, "partial")
            self.assertEqual((loaded[0].flux, saved.iloc[0].Equation), (1, source.iloc[0].Equation))
            self.assertEqual(loaded[0].omitted_components, predictions["EnzymeMap"][0].omitted_components)
            with pd.ExcelFile(output, engine="openpyxl") as excel:
                info = pd.read_excel(excel, sheet_name="Mapping_Info")
            self.assertEqual(len(json.loads(info.iloc[0].OmittedComponentsJSON)), 2)
            wb = load_workbook(output)
            fields = [cell.value for cell in wb["Final_Mapping"][1]]
            wb["Final_Mapping"].cell(2, fields.index("MappingCoverage") + 1, "Complete")
            wb.save(output)
            wb.close()
            with self.assertRaisesRegex(ValueError, "cannot omit"):
                load_final_mappings(output, {})

    def test_manual_priority_standalone_handoff_and_mismatch(self):
        source, metabolites, predictions, _, _ = fixture()
        records = [dict(ReactionID="R", MappingDirection="forward", SourceMetabolite="A[c]",
            ProductMetabolite="B[c]", TransferCoefficient=2, SourceSMILES="CCO", ProductSMILES="CCO",
            CarbonAtomMap="[[0,0,0,0],[0,1,0,1]]", Rule="ManualVerified", RuleVersion="1",
            Evidence="Synthetic user record, not scientific evidence", Notes="Fixture only")]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            manual = root / "manual.csv"
            with manual.open("w", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=["Species", "ReactionID", "MappedReactionSMILES"])
                writer.writeheader()
                writer.writerow(dict(Species="Test", ReactionID="R", MappedReactionSMILES=partial_mapped_smiles(records)))
            selected, choices = select_mappings(predictions)
            apply_manual_mappings("Other", selected, choices, metabolites, manual)
            self.assertEqual(choices.iloc[0].SelectedMapper, "EnzymeMap")
            apply_manual_mappings("Test", selected, choices, metabolites, manual)
            partials = apply_rules(selected, choices, metabolites)
            self.assertEqual(choices.iloc[0].SelectedMapper, "Manual")
            self.assertIn("manual.csv", partials.iloc[0].Evidence)
            self.assertEqual(predictions["EnzymeMap"][0].mapped_smiles, MAPPED)
            bad = deepcopy(records)
            bad[0]["TransferCoefficient"] = 3
            with self.assertRaises(ValueError):
                validate_manual_records(selected[0], bad, {"A[c]":"CCO", "B[c]":"CCO"})
            bad = deepcopy(records)
            bad[0]["CarbonAtomMap"] = "[[0,0,0,0],[0,0,0,1]]"
            with self.assertRaises(ValueError):
                validate_manual_records(selected[0], bad, {"A[c]":"CCO", "B[c]":"CCO"})
            with self.assertRaises(ValueError):
                validate_manual_records(selected[0], records, {"A[c]":"OCC", "B[c]":"CCO"})
            path = root / "input.xlsx"
            with pd.ExcelWriter(path, engine="openpyxl") as writer:
                source.to_excel(writer, sheet_name="Nonzero", index=False)
                metabolites.to_excel(writer, sheet_name="Metabolites", index=False)
            finish_analysis("Test", dict(input_file=path, metabolite_sheet="Metabolites"), metabolites,
                predictions, [], 1, list(predictions), root, stage="map", manual_mapping_file=manual)
            manual.unlink()
            loaded, _, _ = load_final_mappings(root / "Test_2_atom_mapping.xlsx", {})
            self.assertEqual(loaded[0].status, "partial")
            self.assertEqual(float(loaded[0].partial_maps[0]["TransferCoefficient"]), 2)
            self.assertIn("manual.csv", loaded[0].partial_maps[0]["Evidence"])

    def test_source_balances_use_original_stoichiometry(self):
        source, metabolites, _, selected, _ = fixture()
        boundaries = pd.DataFrame([["in", "", 2, "=> A[c]"], ["out", "", 2, "B[c] =>"],
            ["red_in", "", 1, "=> ReducedProtein[c]"], ["ox_out", "", 1, "OxidizedProtein[c] =>"]], columns=source.columns)
        tables, _ = solve_source_contributions(selected + prepare_reactions(boundaries, metabolites), metabolites, ["A[c]"])
        row = tables["MetaboliteContributions"].set_index("MetaboliteName").loc["B[c]"]
        self.assertAlmostEqual(row["A[c]"], 1)
        self.assertAlmostEqual(row.Unknown, 0)
        self.assertEqual((row.Inflow, row.Outflow), (2, 2))

if __name__ == "__main__":
    unittest.main()