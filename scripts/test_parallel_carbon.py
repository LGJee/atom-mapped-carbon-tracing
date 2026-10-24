"""Small deterministic mapper fixture and serial/parallel EMU regression."""
import hashlib
import sys
import tempfile
from pathlib import Path
import unittest
from unittest.mock import patch
import pandas as pd
from openpyxl import Workbook
from carbon_flow_tools.pipeline import run_analyses, run_carbon_flow


def load_model(options):
    return None


def map_batch(model, smiles):
    assert all(text == "CCO>>CCO" for text in smiles)
    return [{"mapped_rxn": "[CH3:1][CH2:2][OH:3]>>[CH3:1][CH2:2][OH:3]", "confidence": 0.8} for _ in smiles]


def make_input(path, product):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Nonzero"
    sheet.append(["ReactionID", "ReactionName", "Equation", "Flux"])
    sheet.append(["r1", "transport", "A[e] -> A[c]", 1])
    sheet.append(["r2", "reaction", "A[c] -> " + product, 1])
    sheet = workbook.create_sheet("Metabolites")
    sheet.append(["MetaboliteID", "MetaboliteName", "SMILES_Yeast9", "SMILES_Manual"])
    for i, name in enumerate(("A[e]", "A[c]", product)):
        sheet.append(["s_" + str(i), name, "CCO", ""])
    workbook.save(path)
    workbook.close()


def prepare_fixture(root, name="Test"):
    source = root / (name + "_1_input.xlsx")
    make_input(source, "B[c]")
    from openpyxl import load_workbook
    wb = load_workbook(source)
    wb["Nonzero"].append(["in", "Import", " => A[e]", 1])
    wb["Nonzero"].append(["out", "Drain", "B[c] => ", 1])
    wb.save(source)
    wb.close()
    job = dict(input_file=source, reaction_sheet="Nonzero", metabolite_sheet="Metabolites",
               carbon_sources=["A[e]"])
    config = {"ExampleMapper": dict(module="scripts.test_parallel_carbon", python=sys.executable, options={"test": True})}
    run_analyses({name: job}, 1, 800, root, ["ExampleMapper"], config, stage="map", max_mapper_processes=1)
    return root / (name + "_2_atom_mapping.xlsx")


class ParallelCarbonTests(unittest.TestCase):
    def test_serial_parallel_source_fractions_are_identical(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            mapping = prepare_fixture(root)
            digest = hashlib.sha256(mapping.read_bytes()).hexdigest()
            jobs = {name: dict(mapping_file=mapping, carbon_sources=["A[e]"], min_available_memory_gb=.01)
                    for name in ("One", "Two")}
            for slots in (1, 2):
                with patch("carbon_flow_tools.pipeline.run_mapper", side_effect=AssertionError("Mapper called in step 3")):
                    statuses = run_carbon_flow(jobs, root / str(slots), max_processes=slots, memory_gb=.25)
                self.assertTrue(all(s["emu_completed"] for s in statuses))
            for name in jobs:
                for sheet in ("EMU_Metabolites", "EMU_BoundaryInputs"):
                    a = pd.read_excel(root / "1" / (name + "_3_carbon_flow.xlsx"), sheet_name=sheet, engine="openpyxl")
                    b = pd.read_excel(root / "2" / (name + "_3_carbon_flow.xlsx"), sheet_name=sheet, engine="openpyxl")
                    pd.testing.assert_frame_equal(a, b)
            self.assertEqual(digest, hashlib.sha256(mapping.read_bytes()).hexdigest())


def main():
    unittest.main(module=__name__)


if __name__ == "__main__":
    main()
