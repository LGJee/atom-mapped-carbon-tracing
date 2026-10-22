"""Partial graphs retain explicit carbon evidence without inventing variable groups."""
from copy import deepcopy
import csv
import hashlib
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import pandas as pd
from rdkit import Chem
from openpyxl import load_workbook

from carbon_flow_tools.carbon_mapping_rules import (
    _REFERENCES, STRUCTURE_LIMITS, apply_rules, apply_manual_mappings, carbon_pairs, carbon_presence,
    carbon_structure_issues, repeat_formula, rule_transfers, structure_kind)
from carbon_flow_tools.carbon_trace import initialize_sources, prepare_mappings, map_one_reaction, propagate_rounds
from carbon_flow_tools.consensus import select_mappings, mapping_signatures
from carbon_flow_tools.pipeline import finish_analysis, load_final_mappings
from carbon_flow_tools.reactions import prepare_reactions
from carbon_flow_tools.smiles_completion import complete_smiles, load_manual
from carbon_flow_tools.model_structures import discover_annotations
from carbon_flow_tools.source_contributions import solve_source_contributions

ACP = "*SC(=O)CCCC(=O)OC"
PC = "*C(=O)OCC(COP(=O)([O-])OCC[N+](C)(C)C)OC(*)=O"
POLYMER = "C20H37O4P(C5H8)n"


def tables(structures, equation, flux=1):
    metabolites = pd.DataFrame([[name, name, structure] for name, structure in structures],
        columns=["MetaboliteID", "MetaboliteName", "SMILES_Manual"])
    reactions = pd.DataFrame([["R", "Fixture", flux, equation]],
        columns=["ReactionID", "ReactionName", "Flux", "Equation"])
    return reactions, metabolites


def marked(smiles):
    mol = Chem.MolFromSmiles(smiles)
    for atom in mol.GetAtoms():
        atom.SetAtomMapNum(atom.GetIdx() + 1)
    return Chem.MolToSmiles(mol)


class PartialStructureTests(unittest.TestCase):
    def test_classification_preserves_normal_syntax(self):
        for text in ["n1ccccc1", "c1cc[nH]c1", "[NH4+]", "[13CH3][C@@H](O)C(=O)[O-]", "[Fe+2].CC(=O)[O-]"]:
            self.assertEqual(structure_kind(text), "complete")
        for text in [ACP, PC, "[H]OC(=O)C([H])([H])C(=O)S[*]", "*", "[*]"]:
            self.assertEqual(structure_kind(text), "partial")
        self.assertEqual(structure_kind("C1CC"), "invalid")
        self.assertEqual(structure_kind(""), "missing")
        self.assertEqual(structure_kind(POLYMER), "polymer")
        self.assertEqual(structure_kind("CC", POLYMER), "polymer")
        self.assertEqual(structure_kind("", POLYMER), "polymer")
        for formula in ["Mn", "Zn", "Sn", "Sm", "C5H5N"]:
            self.assertFalse(repeat_formula(formula))
        self.assertFalse(repeat_formula("n1ccccc1"))
        self.assertTrue(repeat_formula("(C6H10O5)n"))
        self.assertTrue(carbon_presence(SimpleNamespace(MolecularFormula=POLYMER)))
        self.assertIsNone(carbon_presence(SimpleNamespace(SMILES_Manual="*")))
        self.assertTrue(carbon_presence(SimpleNamespace(SMILES_Manual=ACP)))

    def test_manual_scope_legal_wildcards_and_real_errors(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "manual.csv"
            with path.open("w", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["Species", "MetaboliteName", "SMILES"])
                writer.writerows([["Escherichia_coli", "ACP", ACP],
                    ["Other_species", "Broken", "C1CC"], ["Escherichia_coli", "Unused", "C1CC"]])
            before = path.read_bytes()
            self.assertEqual(load_manual(path, species="Escherichia_coli", names={"ACP"})[0]["SMILES"], ACP)
            table = pd.DataFrame([["x", "ACP[c]", "C11H20O5S", -5]],
                columns=["MetaboliteID", "MetaboliteName", "MolecularFormula", "Charge"])
            enriched = complete_smiles(table, "Escherichia_coli_Test", dict(
                lookup=False, manual_file=path, cache_file=Path(folder)/"cache.json", timeout=1, min_interval=0))
            self.assertEqual(enriched.iloc[0].SMILES_Manual, ACP)
            self.assertEqual(enriched.iloc[0].StructureType, "partial")
            self.assertTrue(enriched.iloc[0].StructureReviewRequired)
            self.assertIn("wildcard", enriched.iloc[0].StructureStatus)
            with self.assertRaisesRegex(ValueError, "Invalid manual structure"):
                load_manual(path, species="Escherichia_coli", names={"Unused"})
            self.assertEqual(path.read_bytes(), before)

    def test_actual_manual_table_is_scoped_without_editing(self):
        path = Path(__file__).resolve().parents[1] / "data/manual_smiles.csv"
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        # The actual table has unrelated invalid chemistry; it must not block this compound.
        rows = load_manual(path, species="Escherichia_coli", names={"Glutaryl-[acyl-carrier protein] methyl ester"})
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["SMILES"], ACP)
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)

    def test_companion_annotations_retain_partial_structures(self):
        with tempfile.TemporaryDirectory() as folder:
            model = Path(folder)/"model.json"
            model.write_text(json.dumps(dict(metabolites=[dict(id="a", name="ACP",
                compartment="c", formula="C10H20O5S", annotation=dict(smiles=ACP))])))
            table = pd.DataFrame([["a", "ACP[c]", "C10H20O5S"]],
                columns=["MetaboliteID", "MetaboliteName", "MolecularFormula"])
            found, _ = discover_annotations(dict(path=str(model),
                sha256=hashlib.sha256(model.read_bytes()).hexdigest()), table)
            self.assertEqual(found["a"]["smiles"], ACP)
            self.assertEqual(structure_kind(found["a"]["smiles"]), "partial")

    def test_repeats_are_not_unit_graphs_or_aromatic_nitrogen(self):
        source, metabolites = tables([("ATP[c]", _REFERENCES["ATP"]), ("ADP[c]", _REFERENCES["ADP"]),
            ("polymer[c]", "CC"), ("pool[c]", "")], "55.3 ATP[c] + polymer[c] => 55.3 ADP[c] + pool[c]")
        metabolites["MolecularFormula"] = ["", "", POLYMER, ""]
        reaction = prepare_reactions(source, metabolites)[0]
        self.assertEqual(reaction.equation, source.iloc[0].Equation)
        self.assertEqual(reaction.omitted_components[0]["Reason"], STRUCTURE_LIMITS["polymer"])
        self.assertNotIn("polymer[c]", reaction.mapping_equation)
        selected, choices = select_mappings({"LocalMapper": [reaction]})
        records = apply_rules(selected, choices, metabolites)
        self.assertEqual(records.TransferCoefficient.tolist(), [55.3])
        self.assertEqual(choices.iloc[0].MappingCoverage, "Partial")
        self.assertIn("not set to 1", choices.iloc[0].SelectionReason)
        self.assertTrue(choices.iloc[0].ReviewRequired)
        self.assertEqual(carbon_structure_issues(metabolites), {})
        for text in ["n1ccccc1", "c1cc[nH]c1"]:
            a, m = tables([("A[c]", text), ("B[c]", text)], "A[c] => B[c]")
            self.assertFalse(prepare_reactions(a, m)[0].omitted_components)

    def test_no_variable_group_identity_or_chemical_change_is_invented(self):
        for text in [ACP, PC, "*CC", "*O"]:
            self.assertIsNone(carbon_pairs(text, text))
            source, metabolites = tables([("A[c]", text), ("B[c]", text)], "A[c] => B[c]")
            r = prepare_reactions(source, metabolites)[0]
            self.assertFalse(r.is_transport)
            self.assertFalse(rule_transfers(r, dict(zip(metabolites.MetaboliteName, metabolites.SMILES_Manual))))
            selected, choices = select_mappings({"LocalMapper": [r]})
            self.assertTrue(apply_rules(selected, choices, metabolites).empty)
            self.assertEqual(choices.iloc[0].MappingCoverage, "Unresolved")
            self.assertIn("wildcard", choices.iloc[0].SelectionReason)
        source, metabolites = tables([("ACP[c]", ACP), ("ACP[e]", "*SC(=O)CCC(=O)OC")], "ACP[c] => ACP[e]")
        r = prepare_reactions(source, metabolites)[0]
        self.assertFalse(rule_transfers(r, dict(zip(metabolites.MetaboliteName, metabolites.SMILES_Manual))))

    def test_native_partial_priority_clears_dummy_maps_and_roundtrips(self):
        source, metabolites = tables([("A[c]", ACP), ("B[c]", ACP)], "A[c] => B[c]")
        original = prepare_reactions(source, metabolites)
        predictions = {name: deepcopy(original) for name in ["LocalMapper", "RXNMapper", "EnzymeMap"]}
        raw = marked(ACP) + ">>" + marked(ACP)
        for rows in predictions.values():
            rows[0].mapped_smiles, rows[0].status = raw, "mapped"
        with self.assertRaisesRegex(ValueError, "Partial structure"):
            mapping_signatures(predictions["EnzymeMap"][0])
        selected, choices = select_mappings(predictions)
        self.assertEqual(choices.iloc[0].SelectedMapper, "EnzymeMap")
        self.assertTrue(choices.iloc[0].ReviewRequired)
        self.assertEqual(selected[0].status, "partial")
        for atom in Chem.MolFromSmiles(selected[0].mapped_smiles.split(">>")[0]).GetAtoms():
            if atom.GetAtomicNum() == 0:
                self.assertEqual(atom.GetAtomMapNum(), 0)
        self.assertEqual(predictions["EnzymeMap"][0].mapped_smiles, raw)
        self.check_roundtrip(source, metabolites, predictions)

    def check_roundtrip(self, source, metabolites, predictions):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = root / "input.xlsx"
            with pd.ExcelWriter(path, engine="openpyxl") as writer:
                source.to_excel(writer, sheet_name="Nonzero", index=False)
                metabolites.to_excel(writer, sheet_name="Metabolites", index=False)
            finish_analysis("Test", dict(input_file=path, metabolite_sheet="Metabolites"), metabolites,
                predictions, [], 1, list(predictions), root, stage="map")
            output = root / "Test_2_atom_mapping.xlsx"
            loaded, saved, restored = load_final_mappings(output, {})
            self.assertEqual(saved.iloc[0].Equation, source.iloc[0].Equation)
            self.assertEqual(loaded[0].flux, source.iloc[0].Flux)
            self.assertEqual(loaded[0].status, "partial")
            self.assertEqual(restored.SMILES_Manual.tolist(), metabolites.SMILES_Manual.tolist())
            with pd.ExcelFile(output, engine="openpyxl") as excel:
                self.assertNotIn("Partial_Mapping", excel.sheet_names)
                final = pd.read_excel(excel, sheet_name="Final_Mapping")
                self.assertEqual(final.iloc[0].MappingCoverage, "Partial")
                self.assertTrue(final.iloc[0].ReviewRequired)
            prepared = prepare_mappings(loaded)
            donor = next(iter(prepared[0][1][0]))[0].metabolite
            molecule = Chem.MolFromSmiles(dict(zip(metabolites.MetaboliteName, metabolites.SMILES_Manual))[donor])
            assignment = tuple((0, atom.GetIdx(), "C"+str(i+1)) for i,atom in enumerate(
                atom for atom in molecule.GetAtoms() if atom.GetAtomicNum() == 6))
            _, rows, _ = map_one_reaction(*prepared[0], {donor: {1: assignment}}, 1, {}, 2,
                dict(zip(metabolites.MetaboliteName, metabolites.SMILES_Manual)))
            self.assertTrue(rows)
            self.assertTrue(rows[0]["CarbonLabels"])
            # A false completeness claim is rejected at the step-3 trust boundary.
            wb = load_workbook(output)
            fields = [cell.value for cell in wb["Final_Mapping"][1]]
            wb["Final_Mapping"].cell(2, fields.index("MappingCoverage")+1, "Complete")
            wb.save(output)
            wb.close()
            with self.assertRaisesRegex(ValueError, "Partial structure|completely aligned|Mapped structures"):
                load_final_mappings(output, {})

    def test_manual_partial_does_not_assert_dummy_identity(self):
        source, metabolites = tables([("A[c]", ACP), ("B[c]", ACP)], "A[c] => B[c]")
        selected, choices = select_mappings({"LocalMapper": prepare_reactions(source, metabolites)})
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/"manual.csv"
            raw = marked(ACP) + ">>" + marked(ACP)
            with path.open("w", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["Species", "ReactionID", "MappedReactionSMILES"])
                writer.writerow(["Test", "R", raw])
            before = path.read_bytes()
            apply_manual_mappings("Test", selected, choices, metabolites, path)
            apply_rules(selected, choices, metabolites)
            self.assertEqual(path.read_bytes(), before)
            self.assertEqual(choices.iloc[0].SelectedMapper, "Manual")
            self.assertEqual(choices.iloc[0].MappingCoverage, "Partial")
            self.assertTrue(choices.iloc[0].ReviewRequired)
            for side in selected[0].mapped_smiles.split(">>"):
                self.assertTrue(all(a.GetAtomMapNum() == 0 for a in Chem.MolFromSmiles(side).GetAtoms()
                    if a.GetAtomicNum() != 6))

    def test_unknown_to_carbon_is_not_an_atom_correspondence(self):
        source, metabolites = tables([("A[c]", "*CO"), ("B[c]", "CCO")], "A[c] => B[c]")
        r = prepare_reactions(source, metabolites)[0]
        r.mapped_smiles = "[*:1][CH2:2]O>>[CH3:1][CH2:2]O"
        r.status = "mapped"
        selected, choices = select_mappings({"RXNMapper": [r]})
        self.assertEqual(selected[0].status, "partial")
        mol = Chem.MolFromSmiles(selected[0].mapped_smiles.split(">>")[1])
        mapped_carbons = [atom for atom in mol.GetAtoms() if atom.GetAtomicNum() == 6 and atom.GetAtomMapNum()]
        self.assertEqual(len(mapped_carbons), 1)
        self.assertTrue(choices.iloc[0].ReviewRequired)

    def test_partial_transport_uses_only_invariant_known_carbons(self):
        for text in [ACP, PC]:
            source, metabolites = tables([("Same[c]", text), ("Same[m]", text)], "2 Same[c] => 2 Same[m]", -1)
            r = prepare_reactions(source, metabolites)[0]
            self.assertFalse(r.is_transport)
            selected, choices = select_mappings({"LocalMapper": [deepcopy(r)]})
            records = apply_rules(selected, choices, metabolites)
            self.assertEqual(choices.iloc[0].SelectedMapper, "CarbonRules")
            self.assertEqual(records.iloc[0].Rule, "partial_identity_transport")
            self.assertEqual(records.iloc[0].SourceMetabolite, "Same[m]")
            self.assertEqual(records.iloc[0].TransferCoefficient, 2)
            for a, i, b, j in json.loads(records.iloc[0].CarbonAtomMap):
                self.assertEqual(Chem.MolFromSmiles(text).GetAtomWithIdx(i).GetAtomicNum(), 6)
                self.assertEqual(Chem.MolFromSmiles(text).GetAtomWithIdx(j).GetAtomicNum(), 6)
            self.check_roundtrip(source, metabolites, {"LocalMapper": [r]})

    def test_source_fraction_scope_and_source_validation(self):
        source, metabolites = tables([("Known[c]", "CCO"), ("Partial[c]", "*CCO"),
            ("Polymer[c]", "CC")], "Known[c] => Partial[c]")
        metabolites["MolecularFormula"] = ["C2H6O", "C10H20O", POLYMER]
        boundaries = pd.DataFrame([["in", "", 1, "=> Known[c]"], ["out", "", 1, "Partial[c] =>"]], columns=source.columns)
        reactions = prepare_reactions(pd.concat([source, boundaries], ignore_index=True), metabolites)
        reactions[0].mapped_smiles = "[CH3:1][CH2:2]O>>*[CH2:1][CH2:2]O"
        reactions[0].status = "partial"
        tables_out, info = solve_source_contributions(reactions, metabolites, ["Known[c]"])
        rows = tables_out["MetaboliteContributions"].set_index("MetaboliteName")
        self.assertTrue(pd.isna(rows.loc["Partial[c]", "CarbonCount"]))
        self.assertEqual(rows.loc["Partial[c]", "KnownCarbonCount"], 2)
        self.assertIn("unknown region not quantified", rows.loc["Partial[c]", "ContributionScope"])
        self.assertAlmostEqual(rows.loc["Partial[c]", "Known[c]"], 1)
        self.assertTrue(pd.isna(rows.loc["Polymer[c]", "Known[c]"]))
        self.assertTrue(pd.isna(rows.loc["Polymer[c]", "CarbonCount"]))
        for name, text in [("Partial[c]", "*CCO"), ("Polymer[c]", POLYMER)]:
            with self.assertRaisesRegex(ValueError, "complete, explicit"):
                initialize_sources([name], {name:text})
            with self.assertRaisesRegex(ValueError, "complete, explicit"):
                solve_source_contributions(reactions, metabolites, [name])
        with self.assertRaisesRegex(ValueError, "complete, explicit"):
            next(propagate_rounds("Test", reactions, metabolites, 1, ["Polymer[c]"]))


if __name__ == "__main__":
    unittest.main()
