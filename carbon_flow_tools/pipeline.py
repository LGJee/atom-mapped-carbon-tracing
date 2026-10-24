"""Prediction, consensus selection and fixed-flux carbon-source balances."""

from copy import deepcopy
from datetime import datetime
from hashlib import sha256
import json
from pathlib import Path

from .console import detail
from .task_scheduler import run_tasks

import pandas as pd
from openpyxl import load_workbook, Workbook
from rdkit import Chem, RDLogger

from .atom_helpers import mapped_components
from .config import configure, parallel_capacity, available_cpus, MemoryGuard, MemoryPressureError, DiskSpaceLimitError
from .excel_io import (
    clear_formatting,
    copy_input,
    input_sheets,
    write_sheet,
    read_partial_mapping_records,
    write_partial_mapping_records,
    save_output,
    workbook_limit,
    WorkbookSizeLimitError,
)
from .carbon_mapping_rules import apply_rules, apply_manual_mappings, load_partials, carbon_structure_issues, carbon_presence
from .consensus import select_mappings
from .consensus import mapping_signatures, carbon_candidate
from .mapper_runner import run_mapper, validate_mapper_config, group_mapper_jobs
from .reactions import clean, get_smiles, prepare_reactions, parse_side, split_equation
from .smiles_completion import complete_smiles




def read_input(
    path, reaction_sheet, metabolite_sheet, structure_options=None, model=""
):
    guard = MemoryGuard()
    guard.check_file(path, "reading reaction and metabolite tables")
    try:
        with pd.ExcelFile(path, engine="openpyxl") as excel:
            reactions_df = pd.read_excel(excel, sheet_name=reaction_sheet)
            metabolites = pd.read_excel(excel, sheet_name=metabolite_sheet)
    except ValueError as error:
        raise ValueError(
            "Cannot read input tables.\nFile: {}\nSheets: {}, {}\nDetails: {}\n"
            "Action: check the sheet names and workbook, or regenerate it with step 1."
            .format(path, reaction_sheet, metabolite_sheet, error)
        ) from error
    guard.check("validating input tables", force=True)
    required = {
        "reactions": {"ReactionID", "ReactionName", "Equation", "Flux"},
        "metabolites": {"MetaboliteID", "MetaboliteName"},
    }
    for name, table in (("reactions", reactions_df), ("metabolites", metabolites)):
        missing = required[name] - set(table.columns)
        if missing:
            sheet = reaction_sheet if name == "reactions" else metabolite_sheet
            raise ValueError(
                "Missing {} columns: {}\nFile: {}\nSheet: {}\n"
                "Action: restore the required headers or regenerate the workbook with step 1."
                .format(name, ", ".join(sorted(missing)), path, sheet)
            )
        identity = "ReactionID" if name == "reactions" else "MetaboliteID"
        if (
            table[identity].map(clean).eq("").any()
            or table[identity].duplicated().any()
        ):
            invalid = table[identity].map(clean).eq("") | table[identity].duplicated(keep=False)
            examples = ["{} ({})".format(position + 2, clean(value) or "empty")
                        for position, (bad, value) in enumerate(zip(invalid, table[identity])) if bad]
            sheet = reaction_sheet if name == "reactions" else metabolite_sheet
            raise ValueError(
                "{} must be nonempty and unique.\nFile: {}\nSheet: {}\n"
                "Affected Excel rows (first 8): {}\n"
                "Action: fill missing IDs and resolve duplicates."
                .format(identity, path, sheet, ", ".join(examples[:8]))
            )
    if (
        reactions_df.empty
        or not pd.to_numeric(reactions_df.Flux, errors="coerce")
        .map(lambda value: pd.notna(value) and abs(value) < float("inf") and value != 0)
        .all()
    ):
        invalid = ~pd.to_numeric(reactions_df.Flux, errors="coerce").map(
            lambda value: pd.notna(value) and abs(value) < float("inf") and value != 0)
        examples = ["{} ({})".format(position + 2, identity)
                    for position, (bad, identity) in enumerate(zip(invalid, reactions_df.ReactionID)) if bad]
        raise ValueError(
            "Reaction inputs must contain finite nonzero fluxes.\nFile: {}\nSheet: {}\n"
            "Affected Excel rows / ReactionID (first 8): {}\n"
            "Action: export a nonempty nonzero-flux solution; remove invalid or zero-flux rows."
            .format(path, reaction_sheet, ", ".join(examples[:8]) or "reaction table is empty")
        )
    if "SMILES_Manual" not in metabolites:
        metabolites["SMILES_Manual"] = float("nan")
    if not {"SMILES_Model", "SMILES_Yeast9"}.intersection(metabolites.columns):
        metabolites["SMILES_Model"] = float("nan")
    if structure_options is not None:
        try:
            metabolites = complete_smiles(metabolites, model, structure_options)
        except ValueError as error:
            raise ValueError("{}\nInput workbook: {}\nMetabolite sheet: {}\nModel: {}"
                             .format(error, path, metabolite_sheet, model)) from error
    structures = [get_smiles(row) for row in metabolites.itertuples(index=False)]
    if (
        pd.DataFrame({"name": metabolites.MetaboliteName, "smiles": structures})
        .groupby("name")
        .smiles.nunique()
        .gt(1)
        .any()
    ):
        counts = pd.DataFrame({"name": metabolites.MetaboliteName, "smiles": structures}).groupby("name").smiles.nunique()
        conflicts = counts[counts > 1].index.tolist()
        raise ValueError(
            "A compartment-qualified metabolite name has conflicting structures.\n"
            "File: {}\nSheet: {}\nMetabolites (first 8): {}\n"
            "Action: use one verified structure for each compartment-qualified metabolite."
            .format(path, metabolite_sheet, ", ".join(map(str, conflicts[:8])))
        )
    reactions = prepare_reactions(reactions_df, metabolites)
    return reactions, reactions_df, metabolites


def reuse_predictions(path, backend, settings, reactions, source, metabolites, job):
    """Reuse predictions only when inputs, direction and adapter settings match."""
    if not path.is_file():
        return False
    try:
        with pd.ExcelFile(path, engine="openpyxl") as excel:
            old_source = pd.read_excel(excel, sheet_name=job["reaction_sheet"])
            old_metabolites = pd.read_excel(excel, sheet_name=job["metabolite_sheet"])
            metadata = pd.read_excel(
                excel, sheet_name="Mapping_Info", dtype={"ReactionID": str}
            )
            info = metadata.loc[metadata.Mapper == backend]
            if info.empty:
                return False
            expected = dict(
                Mapper=backend, Adapter=settings["module"], Python=settings["python"]
            )
            if any(
                clean(info.iloc[0].get(key)) != clean(value)
                for key, value in expected.items()
            ):
                return False
            if "OptionsJSON" in info:
                if json.loads(clean(info.iloc[0].OptionsJSON)) != settings.get(
                    "options", {}
                ):
                    return False
            elif any(
                clean(info.iloc[0].get(key)) != clean(value)
                for key, value in settings.get("options", {}).items()
            ):
                return False

            pd.testing.assert_frame_equal(
                old_source[list(source.columns)],
                source,
                check_exact=False,
                rtol=1e-12,
                atol=1e-15,
            )
            if not old_metabolites.equals(metabolites):
                return False
            sheet = backend if backend in excel.sheet_names else job["reaction_sheet"]
            table = pd.read_excel(excel, sheet_name=sheet, dtype={"ReactionID": str})
            if "MappingStatus" not in table:
                details = metadata.loc[
                    (metadata.Mapper == backend) & metadata.ReactionID.notna(),
                    ["ReactionID", "MappingStatus", "MappingTemplate"],
                ]
                table = table.merge(
                    details, on="ReactionID", how="left", validate="one_to_one"
                )
                if table.MappingStatus.isna().any():
                    return False
            if table.ReactionID.tolist() != [r.reaction_id for r in reactions]:
                return False
            for reaction, row in zip(reactions, table.itertuples(index=False)):
                if clean(row.ReactionSMILES) != reaction.reaction_smiles:
                    return False
                if clean(row.MappingDirection) != (
                    "forward" if reaction.flux >= 0 else "reverse"
                ):
                    return False
                status = clean(row.MappingStatus)
                if reaction.status == "ready" and not (
                    status
                    in (
                        "mapped",
                        "Mapped fragments cannot be aligned to input metabolites",
                    )
                    or status.startswith(("Invalid mapping: ", backend + " failed: "))
                ):
                    return False
            for reaction, row in zip(reactions, table.itertuples(index=False)):
                if reaction.is_transport:
                    continue
                reaction.mapped_smiles = clean(row.MappedReactionSMILES)
                reaction.confident = (
                    "" if pd.isna(row.MappingConfident) else row.MappingConfident
                )
                if backend == "LocalMapper" and reaction.confident in (0, 1):
                    reaction.confident = bool(reaction.confident)
                reaction.template = clean(row.MappingTemplate)
                if reaction.status == "ready":
                    reaction.status = clean(row.MappingStatus)
        return True
    except (AssertionError, ValueError, KeyError, AttributeError):
        return False


def prediction_table(reactions):
    return pd.DataFrame(
        [
            {
                "ReactionID": r.reaction_id,
                "ReactionName": r.reaction_name,
                "Equation": r.equation,
                "ReactionSMILES": r.reaction_smiles,
                "MappedReactionSMILES": r.mapped_smiles,
                "MappingDirection": "forward" if r.flux >= 0 else "reverse",
                "MappingConfident": r.confident,
            }
            for r in reactions
        ]
    )


def final_mapping_table(reactions, selection):
    choices = selection.set_index("ReactionID")
    columns = [
        "ReactionID",
        "ReactionName",
        "Flux",
        "SelectedMapper",
        "ReviewRequired",
        "AIGenerated",
        "MappedReactionSMILES",
        "MappingCoverage",
        "Equation",
        "MappingDirection",
    ]
    return pd.DataFrame(
        [
            {
                "ReactionID": r.reaction_id,
                "ReactionName": r.reaction_name,
                "Flux": r.flux,
                "Equation": r.equation,
                "MappingDirection": "forward" if r.flux >= 0 else "reverse",
                "SelectedMapper": choices.loc[r.reaction_id, "SelectedMapper"],
                "MappedReactionSMILES": r.mapped_smiles,
                "ReviewRequired": bool(choices.loc[r.reaction_id, "ReviewRequired"]),
                "AIGenerated": False,
                "MappingCoverage": choices.loc[r.reaction_id].get(
                    "MappingCoverage", "Complete" if r.status == "mapped" else "Unresolved"),
            }
            for r in reactions
        ],
        columns=columns,
    )





def run_carbon_flow(jobs, output_dir, max_processes=None, memory_gb=None):
    """Solve independent original/AI mapping versions."""
    if not jobs:
        raise ValueError("Configure at least one model for carbon-source analysis.")
    output_dir = Path(output_dir).resolve()
    tasks, inputs = [], set()
    for name, job in jobs.items():
        workbook_limit(job.get("max_excel_size_mb", 50))
        if not name or any(c in name for c in '/\\:*?"<>|') or name in (".", ".."):
            raise ValueError("Invalid analysis name for a file: " + name)
        sources = job.get("carbon_sources")
        if not sources or len(set(sources)) != len(sources):
            raise ValueError(name + " carbon sources must be nonempty and unique.")
        mappings = job.get("mapping_files", [
            job.get("mapping_file", output_dir / (name + "_2_atom_mapping.xlsx"))])
        if not mappings:
            raise ValueError("Configure at least one mapping workbook for " + name)
        for mapping in mappings:
            mapping = Path(mapping).resolve()
            if not mapping.is_file():
                raise FileNotFoundError("Run step2_atom_mapping.py first: " + str(mapping))
            inputs.add(mapping)
            if "input_file" in job:
                inputs.add(Path(job["input_file"]).resolve())
            suffix = "_AIGenerated" if mapping.stem.endswith("_AIGenerated") else ""
            tasks.append((name, job, mapping, output_dir / (name + "_3_carbon_flow" + suffix + ".xlsx")))
    outputs = [task[-1] for task in tasks]
    if len(set(outputs)) != len(tasks):
        raise ValueError("Mapping workbooks would overwrite the same carbon-flow output.")
    if set(outputs) & inputs:
        raise ValueError("Output cannot overwrite an input or mapping workbook.")
    workers, threads = parallel_capacity(len(tasks), max_processes, memory_gb, task_kind="carbon-source analysis")
    output_dir.mkdir(parents=True, exist_ok=True)
    previous = {p: p.stat().st_mtime_ns if p.exists() else None for p in outputs}
    try:
        saved = run_tasks(trace_with_status, tasks, workers, memory_gb,
                          weights=[task[2].stat().st_size * len(task[1]["carbon_sources"]) for task in tasks],
                          cpu_threads=threads)
    except Exception as error:
        saved = []
        for name, job, mapping, output in tasks:
            status = dict(model=name, method="EMU", mapping_file=str(mapping),
                          carbon_sources=job["carbon_sources"], output_file=None,
                          emu_completed=False, stop_reason="not_completed")
            # Atomic exports distinguish new complete files from historical outputs.
            if output.is_file() and output.stat().st_mtime_ns != previous[output]:
                status.update(emu_completed=True, output_file=str(output), stop_reason="emu_solved")
            if getattr(error, "run_status", {}).get("mapping_file") == str(mapping):
                status.update(error.run_status)
            saved.append(status)
        write_carbon_run_report(output_dir, tasks, saved, str(error))
        raise
    write_carbon_run_report(output_dir, tasks, saved)
    return saved


def trace_with_status(strain, job, mapping_file, output):
    """Return a compact completion record, including the exact mapping input hash."""
    status = dict(model=strain, method="EMU", emu_completed=False,
                  stop_reason="not_completed", output_file=None,
                  mapping_file=str(mapping_file), carbon_sources=job["carbon_sources"])
    try:
        digest = sha256()
        with mapping_file.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        status["mapping_sha256"] = digest.hexdigest()
        trace_mapping_file(strain, dict(job, _run_status=status), mapping_file, output)
    except Exception as error:
        status["stop_reason"] = ("memory_limit" if isinstance(error, (MemoryPressureError, MemoryError))
                                 else "disk_space_limit" if isinstance(error, DiskSpaceLimitError)
                                 else "file_size_limit" if isinstance(error, WorkbookSizeLimitError) else "error")
        error.run_status = status
        raise
    return status


def write_carbon_run_report(output_dir, tasks, statuses, error=None):
    """Record completion and safety failures without confusing old files with new output."""
    output = output_dir / "carbon_flow_run.json"
    report = dict(version=1, method="EMU", finished_at_utc=datetime.utcnow().isoformat() + "Z",
                  status="error" if error else "finished", models=statuses)
    if error:
        report["error"] = error
    temporary = output.with_suffix(".tmp")
    try:
        temporary.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
        temporary.replace(output)
    finally:
        if temporary.exists():
            temporary.unlink()
    detail("Run report: {}".format(output), flush=True)


def trace_mapping_file(strain, job, mapping_file, output):
    """Validate selected correspondences, solve source balances, and atomically export."""
    from .source_contributions import solve_source_contributions
    from .config import DiskGuard
    RDLogger.DisableLog("rdApp.warning")
    guard = MemoryGuard(job.get("min_available_memory_gb", 2))
    guard.check("loading " + strain, force=True)
    label = "{} / {}".format(strain, "AI" if mapping_file.stem.endswith("_AIGenerated") else "Original")
    print(label + ": reading atom mappings.", flush=True)
    reactions, source, metabolites = load_final_mappings(mapping_file, job)
    print(label + ": calculating carbon-source contributions...", flush=True)
    tables, info = solve_source_contributions(reactions, metabolites, job["carbon_sources"],
        allow_inorganic_carbon_recycling=job.get("emu_allow_inorganic_carbon_recycling", True),
        tolerance=job.get("emu_tolerance", 1e-8), memory_guard=guard)
    # Keep validation and structure-scope metadata internally, outside the Excel export.
    contributions = tables["MetaboliteContributions"].drop(
        columns=["FluxResidual", "Status", "StructureType", "KnownCarbonCount",
                 "ContributionScope"], errors="ignore")
    sheets = {"EMU_Metabolites": contributions, "EMU_BoundaryInputs": tables["BoundaryInputs"]}
    status = job.get("_run_status", {})
    info.update(mapping_version="AI" if mapping_file.stem.endswith("_AIGenerated") else "Original")
    status.update(emu=info, emu_solved=True, emu_completed=False)
    guard.check("exporting EMU results", force=True)
    DiskGuard(output.parent).check(force=True)
    workbook = Workbook()
    del workbook[workbook.active.title]
    try:
        for sheet, table in sheets.items():
            write_sheet(workbook, sheet, table, guard)
        save_output(workbook, output, max_bytes=workbook_limit(job.get("max_excel_size_mb", 50)))
    finally:
        workbook.close()
    status.update(emu_completed=True, output_file=str(output), stop_reason="emu_solved")
    print(label + ": complete.\nSaved: " + str(output), flush=True)
    return output


def load_final_mappings(mapping_file, job):
    """Load selected mappings through the carbon-propagation input checks."""
    MemoryGuard(job.get("min_available_memory_gb", 2)).check_file(mapping_file, "loading final atom mappings")
    with pd.ExcelFile(mapping_file, engine="openpyxl") as excel:
        reaction_sheet, metabolite_sheet = input_sheets(excel)
        final = pd.read_excel(
            excel, sheet_name="Final_Mapping", dtype={"ReactionID": str}
        ).fillna("")
        partials = read_partial_mapping_records(excel.book)
    reactions, source, metabolites = read_input(
        mapping_file, reaction_sheet, metabolite_sheet
    )
    # Preserve the original-input check for callers that explicitly supply step 1.
    if "input_file" in job:
        _, original_source, original_metabolites = read_input(
            job["input_file"], job["reaction_sheet"], job["metabolite_sheet"]
        )
        pd.testing.assert_frame_equal(source, original_source, rtol=1e-12, atol=1e-15)
        pd.testing.assert_frame_equal(metabolites, original_metabolites)
    required = {"ReactionID", "MappingDirection", "Equation", "MappedReactionSMILES"}
    if not required <= set(final.columns):
        raise ValueError("Final_Mapping is missing required columns.")
    if not final.ReactionID.is_unique or not set(final.ReactionID) <= {
        r.reaction_id for r in reactions
    }:
        raise ValueError(
            "Final_Mapping contains duplicated or unknown ReactionID values."
        )
    lookup = final.set_index("ReactionID")
    structure_issues = carbon_structure_issues(metabolites)
    carbon_by_name = {clean(row.MetaboliteName): carbon_presence(row)
                      for row in metabolites.itertuples(index=False)}
    for reaction in reactions:
        reaction.status = "Not selected"
        if reaction.reaction_id not in lookup.index:
            continue
        row = lookup.loc[reaction.reaction_id]
        if (
            row.MappingDirection != ("forward" if reaction.flux >= 0 else "reverse")
            or row.Equation != reaction.equation
        ):
            raise ValueError(
                "Final_Mapping input or direction changed; rerun the mapping stage: {} ({!r} vs {!r})".format(reaction.reaction_id, row.Equation, reaction.equation)
            )
        coverage = row.get("MappingCoverage", "Complete" if row.MappedReactionSMILES else "Unresolved")
        if coverage not in ("Complete", "Partial", "Unresolved", "Boundary", "NoCarbon"):
            raise ValueError("Unknown MappingCoverage: " + str(coverage))
        reaction.mapped_smiles = row.MappedReactionSMILES
        inconsistent = [c.metabolite for c in reaction.left + reaction.right
                        if c.metabolite in structure_issues]
        if coverage in ("Complete", "Partial") and reaction.mapped_smiles and inconsistent:
            raise ValueError("Final mapping uses carbon-inconsistent structures: "
                             + ", ".join(sorted(set(inconsistent)))
                             + "; review structures and rerun step 2 selection")
        if coverage != "Complete":
            if coverage == "NoCarbon":
                names = [name for side in split_equation(reaction.equation) for _, name in parse_side(side)]
                if not names or any(carbon_by_name.get(name) is not False for name in names):
                    raise ValueError("NoCarbon coverage conflicts with saved inputs: " + reaction.reaction_id)
            if coverage == "Boundary" and all(split_equation(reaction.equation)):
                raise ValueError("Boundary coverage requires an open reaction side: " + reaction.reaction_id)
            if coverage in ("Unresolved", "Boundary", "NoCarbon"):
                # Retained failed candidates are audit data, never usable atom transfers.
                reaction.mapped_smiles = ""
                continue
            if coverage == "Partial" and reaction.reaction_id in set(partials.get("ReactionID", [])):
                continue  # Reduced rule subreactions are checked against their exact records below.
            if coverage == "Partial" and reaction.mapped_smiles:
                checked = carbon_candidate(reaction)
                if ((checked.status != "partial" and not reaction.omitted_components)
                        or checked.mapped_smiles != reaction.mapped_smiles):
                    raise ValueError("Saved partial mapper correspondence must be a validated carbon-only view")
                reaction.status = "partial"
            elif reaction.mapped_smiles:
                raise ValueError("Unresolved or boundary coverage cannot contain a mapped reaction")
            continue
        if reaction.omitted_components:
            raise ValueError("Complete coverage cannot omit original reaction participants: " + reaction.reaction_id)
        if not reaction.mapped_smiles:
            raise ValueError("Complete coverage requires a full mapped reaction")
        mapping_signatures(reaction)
        reaction.status = "mapped"
    load_partials(reactions, final, partials, metabolites)
    return reactions, source, metabolites


def run_analyses(
    jobs,
    batch_size,
    max_reaction_atoms,
    output_dir,
    backends,
    mapper_configs,
    refresh=False,
    max_mapper_processes=None,
    mapper_memory_gb=6,
    max_carbon_processes=None,
    carbon_memory_gb=None,
    stage="all",
    structure_options=None,
    manual_mapping_file=None,
):
    """Predict independently, select one mapping per reaction and propagate."""
    if stage == "trace":
        return run_carbon_flow(
            jobs, output_dir, max_carbon_processes, carbon_memory_gb
        )
    if not backends or len(set(backends)) != len(backends):
        raise ValueError("Mapper list must be nonempty and unique.")
    if stage not in ("all", "map", "trace"):
        raise ValueError("Stage must be all, map or trace.")
    for backend in backends:
        if (
            backend
            in (
                "Mapping_Selection",
                "Final_Mapping",
                "Partial_Mapping",
                "Summary",
                "Mapping_Info",
                "Identity",
            )
            or len(backend) > 31
            or any(c in backend for c in "[]")
        ):
            raise ValueError(
                "Mapper names must not conflict with output sheets and must fit Excel limits."
            )
        if stage != "trace":
            validate_mapper_config(backend, mapper_configs[backend])
    configure(batch_size, max_reaction_atoms)
    for name, job in jobs.items():
        if not name or any(c in name for c in '/\\:*?"<>|') or name in (".", ".."):
            raise ValueError(f"Invalid analysis name for a file: {name}")
        if not job["input_file"].is_file():
            raise FileNotFoundError(job["input_file"])
        if stage != "map" and (
            not job["carbon_sources"]
            or len(set(job["carbon_sources"])) != len(job["carbon_sources"])
        ):
            raise ValueError(f"{name} carbon sources must be nonempty and unique.")
        if any(
            backend in (job["reaction_sheet"], job["metabolite_sheet"])
            for backend in backends
        ):
            raise ValueError("Mapper names cannot match input sheet names.")
    input_paths = {job["input_file"].resolve() for job in jobs.values()}
    if any(
        (output_dir / filename).resolve() in input_paths
        for name in jobs
        for filename in (f"{name}_2_atom_mapping.xlsx", f"{name}_3_carbon_flow.xlsx")
    ):
        raise ValueError("Output cannot overwrite any input workbook.")
    output_dir.mkdir(parents=True, exist_ok=True)
    RDLogger.DisableLog("rdApp.warning")
    prepared, pending = {}, []
    for strain, job in jobs.items():
        output = output_dir / f"{strain}_2_atom_mapping.xlsx"
        print(strain + ": checking input and compound structures...", flush=True)
        reactions, source, metabolites = read_input(
            job["input_file"],
            job["reaction_sheet"],
            job["metabolite_sheet"],
            structure_options,
            strain,
        )
        predictions, metadata = {}, []
        for backend in backends:
            settings = mapper_configs[backend]
            candidates = deepcopy(reactions)
            reused = False
            if not refresh:
                reused = reuse_predictions(
                    output, backend, settings, candidates, source, metabolites, job
                )
                if reused:
                    detail(
                        f"{strain} {backend}: inputs and configuration match; reusing predictions.",
                        flush=True,
                    )
            if not reused and any(r.status == "ready" for r in candidates):
                pending.append((strain, backend, candidates, dict(settings,
                    log_file=str(output_dir / "logs" / (strain + "_" + backend + ".log")))))
            predictions[backend] = candidates
            info = dict(
                Mapper=backend,
                Adapter=settings["module"],
                Python=settings["python"],
                ReusedPredictions=reused,
                BatchSize=batch_size,
                MaxReactionAtoms=max_reaction_atoms,
                OptionsJSON=json.dumps(
                    settings.get("options", {}), ensure_ascii=False, sort_keys=True
                ),
                SelectionPolicy="Select carbon agreement supported by at least two tools; otherwise require review; do not compare scores across tools.",
            )
            info.update(settings.get("options", {}))
            info["SourceInputSHA256"] = job.get("input_sha256", "")
            metadata.append(info)
        prepared[strain] = metabolites, predictions, metadata
    grouped = group_mapper_jobs(pending)
    ordinary = [task for task in grouped if task[3]["module"] !=
                "carbon_flow_tools.mappers.enzymemap_adapter"]
    enzyme = [task for task in grouped if task[3]["module"] ==
              "carbon_flow_tools.mappers.enzymemap_adapter"]
    workers, cpu_threads = parallel_capacity(
        len(ordinary), max_mapper_processes, mapper_memory_gb
    )
    reused_count = sum(bool(row["ReusedPredictions"]) for _, _, rows in prepared.values() for row in rows)
    print("Atom mapping: {} tool runs to complete; {} already saved.".format(len(grouped), reused_count), flush=True)
    if pending:
        print("Technical logs: {}".format(output_dir / "logs"), flush=True)
    mapper_tasks = [(*task, cpu_threads, mapper_memory_gb) for task in ordinary]
    run_tasks(run_mapper, mapper_tasks, workers, mapper_memory_gb, processes=False,
              weights=[sum(r.status == "ready" for r in candidates)
                       for _, candidates, _, _ in ordinary])
    # EnzymeMap owns this phase's CPU/memory budget, avoiding nested oversubscription.
    for task in enzyme:
        MemoryGuard().check("starting EnzymeMap", force=True, extra_bytes=1024 ** 3)
        run_mapper(*task, cpu_threads=min(available_cpus(), max_mapper_processes or available_cpus()),
                   memory_gb=None)
    workers, cpu_threads = parallel_capacity(
        len(jobs),
        max_carbon_processes,
        carbon_memory_gb,
        task_kind="mapping selection" if stage == "map" else "carbon propagation",
    )
    tasks = [
        (strain, job, *prepared[strain], backends, output_dir, stage, manual_mapping_file)
        for strain, job in jobs.items()
    ]
    run_tasks(finish_analysis, tasks, workers, carbon_memory_gb,
        weights=[len(next(iter(task[3].values()))) for task in tasks], cpu_threads=cpu_threads)


def finish_analysis(
    strain,
    job,
    metabolites,
    predictions,
    metadata,
    backends,
    output_dir,
    stage="all",
    manual_mapping_file=None,
):
    """Write one selected mapping workbook; optionally run its source balances."""
    RDLogger.DisableLog("rdApp.warning")
    detail(f"{strain}: writing selected atom mappings.", flush=True)
    output = output_dir / f"{strain}_2_atom_mapping.xlsx"
    reactions, selection = select_mappings(predictions, carbon_structure_issues(metabolites))
    apply_manual_mappings(strain, reactions, selection, metabolites, manual_mapping_file, job.get("source_model"))
    partials = apply_rules(reactions, selection, metabolites)
    if "StructureReviewRequired" in metabolites:
        review_names = set(
            metabolites.loc[
                metabolites.StructureReviewRequired.fillna(False).eq(True),
                "MetaboliteName",
            ]
        )
        for index, reaction in enumerate(reactions):
            if any(
                component.metabolite in review_names
                for component in reaction.left + reaction.right
            ):
                selection.at[index, "ReviewRequired"] = True
                selection.at[index, "SelectionReason"] += (
                    "; Automatically selected structure requires review"
                )

    temporary = output.with_suffix(".tmp.xlsx")
    copy_input(job["input_file"], temporary)
    workbook = load_workbook(temporary)
    clear_formatting(workbook)
    write_sheet(workbook, job["metabolite_sheet"], metabolites)
    for backend in backends:
        write_sheet(workbook, backend, prediction_table(predictions[backend]))
    selection_columns = ["ReactionID", "ReactionName", "Flux", "Equation", "MappingDirection",
                         "SelectedMapper", "ReviewRequired", "SelectionReason", "SelectedMappedReactionSMILES"]
    write_sheet(workbook, "Mapping_Selection", selection[selection_columns])
    write_sheet(workbook, "Final_Mapping", final_mapping_table(reactions, selection))
    metadata.extend(
        dict(
            Mapper=backend,
            ReactionID=r.reaction_id,
            MappingStatus=r.status,
            MappingTemplate=r.template,
            MappingInputEquation=r.mapping_equation,
            OmittedComponentsJSON=json.dumps(r.omitted_components, sort_keys=True),
        )
        for backend in backends
        for r in predictions[backend]
    )
    write_sheet(workbook, "Mapping_Info", pd.DataFrame(metadata))
    write_partial_mapping_records(workbook, partials)
    workbook.move_sheet(
        "Mapping_Selection", offset=-workbook.sheetnames.index("Mapping_Selection")
    )
    workbook.move_sheet(
        "Final_Mapping", offset=1 - workbook.sheetnames.index("Final_Mapping")
    )
    save_output(workbook, output)
    chosen = selection.loc[selection.SelectedMapper != ""]
    detail(
        f"{strain} consensus: {int(selection.MappingCoverage.eq('Complete').sum())}/{len(reactions)} complete, "
        f"{int(selection.MappingCoverage.eq('Partial').sum())} partial carbon mappings; "
        f"{int(selection.ReviewRequired.sum())} require review; "
        f"selected tools {chosen.SelectedMapper.value_counts().to_dict()}",
        flush=True,
    )
    print("{}: mappings saved. {}/{} fully mapped; {} partly mapped; {} need review.\nSaved: {}"
          .format(strain, int(selection.MappingCoverage.eq("Complete").sum()),
                  len(reactions), int(selection.MappingCoverage.eq("Partial").sum()),
                  int(selection.ReviewRequired.sum()), output), flush=True)
    workbook.close()
    if stage != "map":
        trace_mapping_file(strain, job, output, output_dir / f"{strain}_3_carbon_flow.xlsx")
