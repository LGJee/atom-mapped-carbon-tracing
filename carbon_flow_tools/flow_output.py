"""Stage propagation rows on disk and publish only complete, bounded workbooks."""

from contextlib import suppress
from pathlib import Path
import sqlite3
import tempfile

import pandas as pd
from openpyxl import Workbook

from .carbon_trace import COLUMNS
from .console import Progress
from .config import DiskGuard
from .propagation_state import create_state_tables
from .atom_helpers import assignment_from_labeled_smiles
from .excel_io import WorkbookSizeLimitError, format_columns


class WorksheetRowLimitError(RuntimeError):
    """The candidate round cannot fit in one worksheet without truncation."""


class FlowOutput:
    MAX_DATA_ROWS = 1048575

    def __init__(self, output, max_bytes, batch_rows, summaries, memory_guard, detail_rounds=None, extra_sheets=None):
        if type(batch_rows) is not int or batch_rows < 1:
            raise ValueError("Output batch rows must be a positive integer.")
        if detail_rounds is not None and (type(detail_rounds) is not int or detail_rounds < 0):
            raise ValueError("Detail rounds must be a nonnegative integer or None")
        self.detail_rounds, self.extra_sheets = detail_rounds, extra_sheets or {}
        self.output = Path(output)
        self.output.parent.mkdir(parents=True, exist_ok=True)
        self.disk_guard = DiskGuard(self.output.parent)
        self.disk_guard.check(force=True)
        self.directory = tempfile.TemporaryDirectory(prefix=".carbon-flow-", dir=str(self.output.parent))
        self.candidate = Path(self.directory.name) / "candidate.xlsx"
        self.db = sqlite3.connect(str(Path(self.directory.name) / "rows.sqlite"))
        self.db.execute("PRAGMA cache_size=-2048")
        self.db.execute("PRAGMA temp_store=FILE")
        self.db.execute("PRAGMA journal_mode=OFF")
        columns = ", ".join('"{}" {}'.format(c, "INTEGER" if c in ("Step", "ProductionID") else "TEXT") for c in COLUMNS)
        self.fields = ", ".join('"{}"'.format(c) for c in COLUMNS)
        self.db.execute("CREATE TABLE records (seq INTEGER PRIMARY KEY, round_no INTEGER, phase INTEGER, layer INTEGER, "
                        + columns + ", UNIQUE(round_no, " + self.fields + "))")
        self.insert = "INSERT OR IGNORE INTO records (round_no, phase, layer, " + self.fields + ") VALUES (" + ",".join("?" for _ in range(len(COLUMNS) + 3)) + ")"
        self.db.execute("CREATE TABLE round_states (name TEXT, id INTEGER PRIMARY KEY, marked TEXT, labels TEXT, step INTEGER)")
        self.states, self.catalog, self.first_steps = create_state_tables(self.db, memory_guard)
        self.compact, self.compact_buffer, self.visits = False, {}, 0
        self.max_bytes, self.batch_rows = max_bytes, batch_rows
        self.summaries, self.memory_guard = summaries, memory_guard
        self.buffer, self.buffer_bytes = [], 0
        self.exclusive_metabolites = None
        self.number, self.phase, self.layer, self.row_count = 0, 0, 0, 0

    def start_round(self, number):
        if self.detail_rounds is not None and number > self.detail_rounds:
            self.db.execute("DELETE FROM records WHERE round_no > ?", (self.detail_rounds,))
            self.db.commit()
        self.db.execute("DELETE FROM round_states")
        self.compact_buffer.clear()
        self.compact = self.detail_rounds is not None and number > self.detail_rounds
        self.number, self.phase, self.layer, self.row_count = number, 0, 0, 0

    def start_layer(self, phase=0):
        self.phase, self.layer = phase, self.layer + 1

    def append_state(self, name, number, marked, labels, step):
        """Keep one minimum-step record per product state, never provenance rows."""
        self.visits += 1
        if self.visits % 256 == 0:
            if self.memory_guard:
                self.memory_guard.check("aggregating compact propagation states")
            self.disk_guard.check()
        previous = self.compact_buffer.get(number)
        if previous is None or step < previous[-1]:
            self.compact_buffer[number] = (name, number, marked, labels, step)
            if previous is None:
                self.buffer_bytes += 256 + 4 * (len(name) + len(marked) + len(labels))
        if len(self.compact_buffer) >= self.batch_rows or self.buffer_bytes >= 8 * 1024 ** 2:
            self.flush()

    def append(self, row):
        if self.compact:
            self.append_state(row["Production"], row["ProductionID"], row["LabeledSMILES"],
                              row["CarbonLabels"], row["Step"])
            return
        self.buffer.append((self.number, self.phase, self.layer) + tuple(row[c] for c in COLUMNS))
        self.buffer_bytes += 512 + sum(4 * len(value) for value in row.values() if isinstance(value, str))
        if len(self.buffer) >= self.batch_rows or self.buffer_bytes >= 8 * 1024 ** 2:
            self.flush()

    def flush(self):
        self.disk_guard.check()
        if self.compact_buffer:
            if self.memory_guard:
                self.memory_guard.check("flushing compact propagation states", force=True)
            self.db.executemany("INSERT INTO round_states VALUES (?,?,?,?,?) "
                "ON CONFLICT(id) DO UPDATE SET step=MIN(step,excluded.step)", self.compact_buffer.values())
            self.db.commit()
            self.compact_buffer.clear()
            self.buffer_bytes = 0
        if not self.buffer:
            return
        if self.memory_guard:
            self.memory_guard.check("flushing carbon-flow batch", force=True)
        before = self.db.total_changes
        self.db.executemany(self.insert, self.buffer)
        self.db.commit()
        self.buffer.clear()
        self.buffer_bytes = 0
        inserted = self.db.total_changes - before
        self.row_count += inserted
        if inserted and (self.detail_rounds is None or self.number <= self.detail_rounds):
            # All exported rounds share one worksheet and therefore one row limit.
            exported_rows = self.db.execute("SELECT COUNT(*) FROM records WHERE round_no<=?", (self.number,)).fetchone()[0]
            if exported_rows > self.MAX_DATA_ROWS:
                raise WorksheetRowLimitError("Carbon_Flow exceeds {} combined worksheet data rows; incomplete round {} discarded.".format(self.MAX_DATA_ROWS, self.number))
            self.checkpoint()

    def merge_states(self, states, first_steps, terminal_names, smiles):
        self.flush()
        rows = (self.db.execute('SELECT name,id,marked,step FROM round_states ORDER BY id') if self.compact else
                self.db.execute('SELECT Production, ProductionID, LabeledSMILES, MIN(Step) FROM records '
                    'WHERE round_no=? AND layer=? GROUP BY Production, ProductionID', (self.number, self.layer)))
        for index, (name, number, marked, step) in enumerate(rows):
            if self.memory_guard and index % 1024 == 0:
                self.memory_guard.check("merging staged propagation states")
            if name not in terminal_names and number not in states[name]:
                states[name][number] = assignment_from_labeled_smiles(smiles[name], marked)
            key = (name, number)
            first_steps[key] = min(first_steps.get(key, step), step)

    def source_records(self):
        if self.compact:
            return self.db.execute('SELECT name,labels FROM round_states')
        return self.db.execute('SELECT DISTINCT Production, CarbonLabels FROM records WHERE round_no=?', (self.number,))

    def checkpoint(self):
        """Measure a real candidate XLSX; its incomplete round is never published."""
        if self.memory_guard:
            self.memory_guard.check("saving carbon-flow checkpoint", force=True)
        self.disk_guard.check(force=True)
        workbook = Workbook(write_only=True)
        for font in workbook._fonts:
            font.name, font.scheme, font.charset = "Calibri", None, None
        try:
            summary = pd.DataFrame(self.summaries).drop(
                columns=["Method", "MixedSources_LabeledCarbonMetabolites",
                         "InorganicCarbonRecycling"], errors="ignore")
            sheet = workbook.create_sheet("Summary")
            format_columns(sheet, list(summary.columns))
            sheet.append(list(summary.columns))
            for row in summary.itertuples(index=False, name=None):
                sheet.append(list(row))
            last_detail = self.number if self.detail_rounds is None else min(self.number, self.detail_rounds)
            if last_detail:
                sheet = workbook.create_sheet("Carbon_Flow")
                columns = ["Carbon_Flow_Round"] + COLUMNS
                format_columns(sheet, columns)
                sheet.append(columns)
                rows = self.db.execute("SELECT round_no, " + self.fields + " FROM records WHERE round_no<=? "
                                       "ORDER BY round_no, phase, Step, ReactionID, SourceIDs, ProductionID, seq", (last_detail,))
                total = self.db.execute(
                    "SELECT COUNT(*) FROM records WHERE round_no<=?", (last_detail,)
                ).fetchone()[0]
                progress = Progress(self.output.stem + " | preparing Excel rows", total, "rows")
                for index, row in enumerate(rows):
                    if index % 1024 == 0:
                        progress.update(index)
                        self.disk_guard.check()
                    if self.memory_guard and index % 1024 == 0:
                        self.memory_guard.check("streaming staged Excel rows")
                    sheet.append(row)
                progress.update(total)
            for name, table in self.extra_sheets.items():
                sheet = workbook.create_sheet(name)
                format_columns(sheet, list(table.columns))
                sheet.append(list(table.columns))
                for row in table.itertuples(index=False, name=None):
                    sheet.append([None if pd.isna(value) else value for value in row])
            if self.exclusive_metabolites is not None:
                columns = ["Round", "CarbonSource", "MetaboliteID", "MetaboliteName"]
                sheet = workbook.create_sheet("Exclusive_Source_Metabolites")
                format_columns(sheet, columns)
                sheet.append(columns)
                for row in self.exclusive_metabolites:
                    sheet.append(row)
            workbook.save(self.candidate)
        finally:
            # Also release openpyxl's temporary XML files after a failed write.
            for sheet in workbook:
                with suppress(Exception):
                    sheet.close()
                if sheet._writer is not None:
                    with suppress(OSError):
                        sheet._writer.cleanup()
            workbook.close()
        size = self.candidate.stat().st_size
        if size >= self.max_bytes:
            raise WorkbookSizeLimitError(
                "Round {} candidate is {:.3f} MB at a batch checkpoint (limit {:.3f} MB); "
                "incomplete round discarded.".format(self.number, size / 1e6, self.max_bytes / 1e6))

    def complete_round(self):
        self.flush()
        self.checkpoint()
        self.candidate.replace(self.output)

    def close(self):
        self.buffer.clear()
        self.compact_buffer.clear()
        self.catalog.cache.clear()
        self.states.page.cache_clear()
        self.db.close()
        self.directory.cleanup()
