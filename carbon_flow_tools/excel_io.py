import shutil
import math
from copy import copy
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter


def input_sheets(excel):
    """Find the original input tables by their headers, independent of sheet names."""
    reaction_sheets, metabolite_sheets = [], []
    for sheet in excel.sheet_names:
        if sheet in ("Mapping_Selection", "Final_Mapping", "Mapping_Info"):
            continue
        columns = set(pd.read_excel(excel, sheet_name=sheet, nrows=0).columns)
        if {"ReactionID", "ReactionName", "Flux", "Equation"} <= columns:
            reaction_sheets.append(sheet)
        if {"MetaboliteID", "MetaboliteName"} <= columns:
            metabolite_sheets.append(sheet)
    if len(reaction_sheets) != 1 or len(metabolite_sheets) != 1:
        raise ValueError(
            "Workbook must contain exactly one reaction table and one metabolite table; "
            "found reaction sheets {}, metabolite sheets {}.".format(
                reaction_sheets, metabolite_sheets
            )
        )
    return reaction_sheets[0], metabolite_sheets[0]


def format_equation_column(worksheet):
    """Limit equation width and wrap long equations."""
    for header in worksheet[1]:
        if header.value != "Equation":
            continue
        dimension = worksheet.column_dimensions[header.column_letter]
        dimension.width = 50
        dimension.bestFit = False
        for row in worksheet.iter_rows(min_col=header.column, max_col=header.column):
            cell = row[0]
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            cell.alignment = alignment
        if worksheet.protection.sheet:
            worksheet.protection.formatColumns = False
        break


def clear_formatting(workbook):
    """Keep plain output tables with wrapped equations."""
    for worksheet in workbook:
        for row in worksheet:
            for cell in row:
                cell._style = None
        for dimension in worksheet.row_dimensions.values():
            dimension.height = None
            dimension._style = None
        for dimension in worksheet.column_dimensions.values():
            dimension._style = None
        worksheet.sheet_format.defaultRowHeight = 15
        worksheet.freeze_panes = None
        worksheet.auto_filter.ref = None
        worksheet.sheet_view.showGridLines = True
        worksheet.sheet_properties.tabColor = None
        for rule in list(worksheet.conditional_formatting):
            del worksheet.conditional_formatting[rule]
        format_columns(worksheet)


def copy_input(source: Path, destination: Path) -> Path:
    """Copy input before analysis; never write results into the input."""
    if source.resolve() == destination.resolve():
        raise ValueError("Output cannot overwrite the input workbook.")
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(str(source), str(destination))
    print(f"Input copied; analysis output: {destination}", flush=True)
    return destination


def write_sheet(
    workbook, name: str, table: pd.DataFrame, memory_guard=None
) -> None:
    if memory_guard:
        memory_guard.check("writing " + name, force=True)
    if name in workbook.sheetnames:
        del workbook[name]
    worksheet = workbook.create_sheet(name)
    worksheet.sheet_state = "hidden" if name == "Mapping_Info" else "visible"
    worksheet.append(list(table.columns))
    for index, row in enumerate(table.itertuples(index=False, name=None)):
        if memory_guard and index % 1024 == 0:
            memory_guard.check("writing " + name)
        worksheet.append([None if pd.isna(value) else value for value in row])
    format_columns(worksheet)


class WorkbookSizeLimitError(RuntimeError):
    """Reject a candidate workbook without replacing the last completed output."""


def workbook_limit(megabytes):
    """Validate the output limit and convert decimal MB to bytes."""
    if (type(megabytes) not in (int, float) or not math.isfinite(megabytes)
            or megabytes * 1e6 < 1):
        raise ValueError("Maximum Excel size must be finite and at least one byte.")
    return int(megabytes * 1e6)


def save_output(workbook, output, max_bytes=None):
    """Keep the previous completed workbook until its replacement is saved."""
    # Use an explicit font rather than the locale-dependent theme font.
    for font in workbook._fonts:
        font.name = "Calibri"
        font.scheme = None
        font.charset = None
    workbook._fonts._dict = {font: i for i, font in enumerate(workbook._fonts)}
    temporary = output.with_suffix(".tmp.xlsx")
    workbook.save(temporary)
    if max_bytes is not None and temporary.stat().st_size > max_bytes:
        size = temporary.stat().st_size
        temporary.unlink()
        raise WorkbookSizeLimitError(
            "Candidate workbook is {:.3f} MB; limit is {:.3f} MB."
            .format(size / 1e6, max_bytes / 1e6)
        )
    try:
        temporary.replace(output)
    except PermissionError:
        raise PermissionError(
            f"Close {output.name} in Excel; completed output is in {temporary.name}."
        )


def format_columns(worksheet, columns=None):
    """Reuse readable plain column widths for copied inputs and generated tables."""
    widths = {
        "ReactionID": 18,
        "MetaboliteID": 18,
        "MetaboliteName": 38,
        "SMILES_Model": 65,
        "SMILES_Yeast9": 65,
        "SMILES_Manual": 65,
        "SMILES_Retrieved": 65,
        "StructureSource": 24,
        "StructureReference": 26,
        "StructureCandidates": 35,
        "StructureReviewRequired": 26,
        "StructureStatus": 50,
        "ReactionName": 42,
        "Flux": 14,
        "Equation": 50,
        "ReactionSMILES": 65,
        "MappedReactionSMILES": 65,
        "Step": 10,
        "Source": 38,
        "SourceIDs": 20,
        "Production": 38,
        "ProductionID": 15,
        "CarbonLabels": 30,
        "LabeledSMILES": 70,
        "SelectedMapper": 18,
        "ReviewRequired": 18,
        "SelectionReason": 60,
        "SupportingMappers": 35,
        "ValidMappers": 35,
        "MappingTemplate": 55,
        "SelectedMappedReactionSMILES": 65,
        "MappingStatus": 40,
        "MappingDirection": 18,
        "HeavyAtomAgreement": 22,
    }
    headers = columns if columns is not None else [cell.value for cell in worksheet[1]]
    for index, header in enumerate(headers, 1):
        column, letter = str(header), get_column_letter(index)
        width = (
            44
            if column.endswith(("_ScoreType", "_Validation"))
            else widths.get(column, 18)
        )
        if "_LabeledCarbon" in column or column == "TotalCarbonMetabolites":
            width = min(80, max(width, len(column) + 2))
        worksheet.column_dimensions[letter].width = width
    if columns is None:
        format_equation_column(worksheet)
